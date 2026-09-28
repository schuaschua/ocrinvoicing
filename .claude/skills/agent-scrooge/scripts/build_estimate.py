#!/usr/bin/env python3
# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Build a cost estimate workbook (.xlsx) from an estimate model, and check it for drift.

The model (<name>.estimate.json, in the org config's cost_estimates_folder, default
docs/costing/) lists every environment's cloud line items, the delivery team, other one-off
costs and the gaps, each line citing where it came from. The workbook has:

  Overall          inputs (days of delivery, working days per month, hours per month, years of run
                   after go-live (3-5), yearly increase in run costs, contingency, exchange rate), the
                   delivery team (employees x days x allocation = mandays, priced from the rate card),
                   the run support team (OPEX: mandays per year) and the cost summary gathering every sheet
  Phase Budgets    the Claude API budget for each planning phase still to run (test design, governance ...),
                   against its actual from the ledger once the close-out sheets are in
  Multi-Year       the build, then Year 1..5 (zero past the horizon): revenue, every cost, contingency,
                   net and cumulative net, payback and return on cost
  Revenue          yearly revenue per stream, from the owner or the business case
  Run Costs        recurring yearly costs outside the environments and the support team
  <environment>    one sheet per environment: component, Azure service, SKU, unit price, quantity,
                   usage per month, monthly, daily and annual cost, citation and price source
  Rate Card        role -> day rate, from the project's rate card
  Other Costs      one-off costs (licences, pen test, training ...)
  Sources & Gaps   price basis, assumptions, what is not priced yet, and the stamp of every cited file

Every data range is a named Excel table (tblInputs, tblTeam, tblCostSummary, tblEnv<name>, tblRateCard,
tblOtherCosts, tblPhaseBudgets, and with close_out tblClaudeUsage, tblClaudeUsageDetail, tblClaudePrices, tblTimeline,
tblAIFeedback),
so Power BI and Excel Online can load them by name. <name>.infra_cost.csv, also beside the workbook, has the
infrastructure OPEX after go-live: every priced cloud line, each environment's total and the grand total, as
daily, weekly, monthly and annual cost (daily = monthly x 12 / 365, weekly = daily x 7, annual = monthly x the
environment's run months per year; year-one prices, before any yearly increase or contingency). Formulas get their values when Excel first opens
the file (Power BI reads the values Excel last saved), so <name>.facts.csv is also written beside the
workbook, with every cost line as a plain value for dashboards that read files directly.

Every total is an Excel formula, so changing a blue input cell recalculates the workbook. The
figures are also computed here and printed, so they can be reported without opening Excel.

Model (JSON):
  {"name": "burdeebur", "title": "...", "currency": "AED",           reporting currency
   "cloud_currency": "USD", "fx_rate": 3.6725, "fx_source": "...",    only when prices differ
   "pricing_region": "uaenorth", "price_date": "2026-09-28",
   "assumptions": {"delivery_days": 60, "working_days_per_month": 21, "hours_per_month": 730,
                   "horizon_years": 5, "cost_escalation_pct": 3, "contingency_pct": 10},
   "environments": [{"name": "Dev", "delivery_months": null, "run_months_per_year": 12, "items": [
       {"component": "API", "service": "Azure Kubernetes Service", "sku": "Standard_D4s_v5 node",
        "unit": "1 Hour", "unit_price": 0.235, "quantity": 2, "usage": "hours",
        "cite": "_bmad-output/planning-artifacts/architecture/x/ARCHITECTURE-SPINE.md AD-4",
        "price_source": "Azure Retail Prices API ... meterId ..., fetched 2026-09-28", "note": ""}]}],
   "team": [{"role": "Developer", "employees": 3, "days": null, "allocation": 1, "cite": "..."}],
   "support": [{"role": "Developer", "mandays_per_year": 60, "cite": "owner: Delivery Manager, 2026-09-28"}],
   "run_costs": [{"item": "Support contract", "annual_cost": 12000, "cite": "owner: ..."}],
   "revenue": [{"stream": "Subscriptions", "basis": "1,000 customers x $20/month", "amounts": [240000, 300000, 360000,
                420000, 480000], "cite": "_bmad-output/specs/SPEC.md 7 Business case"}],
   "other_costs": [{"item": "Penetration test", "cost": 40000, "cite": "owner: Tech Lead, 2026-09-28"}],
   "phase_budgets": [{"phase": "test design", "api_budget": 40, "cite": "owner: Delivery Manager, 2026-09-28"}],
   "gaps": ["Prod DR region: the spine does not say whether there is one"],
   "close_out": {"fx_rate": 3.6725, "fx_source": "...", "in_totals": true}}    optional, at the project's end

  close_out adds the Claude Usage, Claude Usage Detail, Claude Prices and Project Timeline sheets from
  the token ledger's claude_usage.csv and project_timeline.csv (token_usage_folder) and
  _bmad/memory/agent-scrooge/claude-prices.json (paths overridable with "ledger", "prices",
  "timeline"), and a Claude API usage line in the cost summary ("in_totals": false shows it as a memo
  below the grand total instead, for Claude Code on a subscription). fx_rate turns USD into the
  reporting currency. When the team has rated the AI at the end of its epics (ai_feedback.py), close_out
  also adds the AI Feedback sheet from _bmad/memory/agent-scrooge/ai-feedback.csv ("feedback").

  cite: "<project-relative file> [<text that appears in it, e.g. AD-4 or a heading>]",
        "owner: <who, date>" for a figure the owner gave, or "assumption: <what and why>".
  usage: units per month per quantity; "hours" means the Hours per month input (always-on).
  null delivery_months / days mean the Overall inputs; run_months_per_year defaults to 12.
  revenue amounts: one per year of the horizon (up to 5), never invented: from the owner or a cited business case.
  phase_budgets: the Claude API budget (reporting currency, API list prices) per planning phase, named as the token
  ledger names it (claude_usage.csv work_item). A control figure beside the epics' and sprints' api_budget: it is
  compared with the actual, never added to a total (the close-out's Claude API usage line already holds the actual).

Rate card: <project>/_bmad/memory/agent-scrooge/rate-card.json (or --rate-card, or "rate_card" in
the model): {"currency": "AED", "updated": "2026-09-28", "source": "...",
             "roles": [{"role": "Developer", "day_rate": 2500, "note": ""}]}

Usage:
    uv run build_estimate.py <project-root> <model.json> [--rate-card FILE] [--out FILE] [--force]
    uv run build_estimate.py <project-root> <model.json> --check     validate only
    uv run build_estimate.py <project-root> <model.json> --drift     cited files changed since the build?

Building writes <name>.xlsx beside the model (or --out) and a "stamp" into the model (build time,
sha256 of every cited file and the rate card, the workbook's save time). It refuses to overwrite a workbook edited after the
last build unless --force. Prints JSON. Exit codes: 0 ok, 1 invalid model or drift found, 2 error.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import tomllib
from datetime import datetime, timedelta
from pathlib import Path

DATA_DIR = Path("_bmad") / "memory" / "agent-scrooge"
RESERVED = {"overall", "rate card", "other costs", "run costs", "revenue", "multi-year", "epics", "sprints", "phase budgets",
            "sources & gaps",
            "claude prices", "claude usage", "claude usage detail", "claude usage daily", "project timeline"}
SPARE_ROWS = 3
DEFAULTS = {"working_days_per_month": 21, "hours_per_month": 730, "horizon_years": 5, "cost_escalation_pct": 0,
            "contingency_pct": 0, "sprint_days": 10}
MAX_YEARS = 5  # the Multi-Year sheet always has five year columns; years past the horizon show zero

# --------------------------------------------------------------------------- inputs


def org_config(root: Path) -> dict:
    """The [modules.org] section of _bmad/config.toml, with _bmad/custom overrides; {} without one."""
    merged: dict = {}
    for path in (root / "_bmad" / "config.toml", root / "_bmad" / "custom" / "config.toml",
                 root / "_bmad" / "custom" / "config.user.toml"):
        if path.is_file():
            with open(path, "rb") as fh:
                merged.update(tomllib.load(fh).get("modules", {}).get("org", {}))
    return merged


def token_usage_folder(root: Path) -> Path:
    value = str(org_config(root).get("token_usage_folder") or "docs/costing/token_usage").replace("{project-root}", str(root))
    path = Path(value)
    return path if path.is_absolute() else root / path


def estimates_folder(root: Path) -> Path:
    value = str(org_config(root).get("cost_estimates_folder") or "docs/costing")
    value = value.replace("{project-root}", str(root))
    path = Path(value)
    return path if path.is_absolute() else root / path


def load_rate_card(root: Path, model: dict, override: Path | None) -> tuple[dict, Path | None]:
    if "rate_card" in model:
        return model["rate_card"], None
    path = override or root / DATA_DIR / "rate-card.json"
    if not path.is_file():
        return {}, path
    return json.loads(path.read_text(encoding="utf-8")), path


def is_number(value, minimum: float = 0, strict: bool = False) -> bool:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return False
    return value > minimum if strict else value >= minimum


def split_cite(cite: str) -> tuple[str, str, str]:
    """(kind, file, ref): kind is file, owner or assumption."""
    cite = cite.strip()
    for kind in ("owner", "assumption"):
        if cite.lower().startswith(kind + ":"):
            return kind, "", cite[len(kind) + 1:].strip()
    file, _, ref = cite.partition(" ")
    return "file", file, ref.strip()


def check_cite(root: Path, cite, where: str, errors: list[str]) -> None:
    if not isinstance(cite, str) or not cite.strip():
        errors.append(f"{where}: no cite (a project file, 'owner: <who, date>' or 'assumption: <why>')")
        return
    kind, file, ref = split_cite(cite)
    if kind != "file":
        if not ref:
            errors.append(f"{where}: '{kind}:' cite says nothing")
        return
    path = (root / file).resolve()
    if not path.is_relative_to(root.resolve()):
        errors.append(f"{where}: cited file is outside the project: {file}")
    elif not path.is_file():
        errors.append(f"{where}: cited file does not exist: {file}")
    elif ref and ref.lower() not in path.read_text(encoding="utf-8", errors="replace").lower():
        errors.append(f"{where}: '{ref}' does not appear in {file}")


def validate(root: Path, model: dict, card: dict) -> list[str]:
    errors: list[str] = []
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]*", str(model.get("name", ""))):
        errors.append("name: letters, digits, '.', '_' and '-' only")
    ccy = model.get("currency")
    if not (isinstance(ccy, str) and re.fullmatch(r"[A-Z]{3}", ccy)):
        errors.append("currency: a three-letter code such as USD or AED")
    cloud_ccy = model.get("cloud_currency", ccy)
    if cloud_ccy != ccy:
        if not is_number(model.get("fx_rate"), strict=True):
            errors.append(f"fx_rate: needed to turn {cloud_ccy} prices into {ccy}")
        if not model.get("fx_source"):
            errors.append("fx_source: say where the exchange rate came from and when")
    a = model.get("assumptions", {})
    if not is_number(a.get("delivery_days"), strict=True):
        errors.append("assumptions.delivery_days: the number of working days of delivery")
    for key in ("working_days_per_month", "hours_per_month"):
        if key in a and not is_number(a[key], strict=True):
            errors.append(f"assumptions.{key}: a positive number")
    for key in ("cost_escalation_pct", "contingency_pct"):
        if key in a and not is_number(a[key]):
            errors.append(f"assumptions.{key}: zero or more")
    if "run_months" in a:
        errors.append("assumptions.run_months: replaced by horizon_years (3 to 5 years of run after go-live)")
    horizon = a.get("horizon_years", DEFAULTS["horizon_years"])
    if horizon not in (3, 4, 5):
        errors.append("assumptions.horizon_years: 3, 4 or 5 years of run after go-live")

    envs = model.get("environments") or []
    if not envs:
        errors.append("environments: at least one")
    seen: set[str] = set()
    for e_i, env in enumerate(envs):
        name = str(env.get("name", ""))
        where = f"environments[{e_i}] '{name}'"
        if not name or len(name) > 31 or re.search(r"[\[\]:*?/\\']", name):
            errors.append(f"{where}: sheet name must be 1-31 characters without []:*?/\\'")
        if name.lower() in RESERVED or name.lower() in seen:
            errors.append(f"{where}: name clashes with another sheet")
        seen.add(name.lower())
        if env.get("delivery_months") is not None and not is_number(env["delivery_months"]):
            errors.append(f"{where}.delivery_months: zero or more, or null for the Overall input")
        if "run_months" in env:
            errors.append(f"{where}.run_months: replaced by run_months_per_year (how many months a year it runs, default 12)")
        per_year = env.get("run_months_per_year", 12)
        if not is_number(per_year) or per_year > 12:
            errors.append(f"{where}.run_months_per_year: 0 to 12")
        if not env.get("items"):
            errors.append(f"{where}: no items (an environment with nothing priced belongs in gaps)")
        for i_i, item in enumerate(env.get("items") or []):
            iw = f"{name} item {i_i + 1} '{item.get('component', '')}'"
            for key in ("component", "service", "unit", "price_source"):
                if not item.get(key):
                    errors.append(f"{iw}: {key} missing")
            if not is_number(item.get("unit_price")):
                errors.append(f"{iw}: unit_price must be a number (an unpriced item belongs in gaps)")
            if not is_number(item.get("quantity"), strict=True):
                errors.append(f"{iw}: quantity must be above zero")
            usage = item.get("usage", "hours" if "hour" in str(item.get("unit", "")).lower() else None)
            if usage != "hours" and not is_number(usage, strict=True):
                errors.append(f"{iw}: usage must be 'hours' or a number above zero")
            check_cite(root, item.get("cite"), iw, errors)

    if card.get("currency") and ccy and card["currency"] != ccy:
        errors.append(f"rate card is in {card['currency']}, the estimate in {ccy}")
    rates = {}
    for r in card.get("roles", []):
        if not r.get("role") or not is_number(r.get("day_rate"), strict=True):
            errors.append(f"rate card: role '{r.get('role', '')}' needs a day_rate above zero")
        rates[str(r.get("role", "")).lower()] = r.get("day_rate")
    for t_i, member in enumerate(model.get("team") or []):
        tw = f"team[{t_i}] '{member.get('role', '')}'"
        if str(member.get("role", "")).lower() not in rates:
            errors.append(f"{tw}: role is not on the rate card")
        if not is_number(member.get("employees"), strict=True):
            errors.append(f"{tw}: employees must be above zero")
        if member.get("days") is not None and not is_number(member["days"], strict=True):
            errors.append(f"{tw}: days must be above zero, or null for the days of delivery")
        alloc = member.get("allocation", 1)
        if not is_number(alloc, strict=True) or alloc > 1:
            errors.append(f"{tw}: allocation is a fraction above 0 and at most 1")
        check_cite(root, member.get("cite"), tw, errors)
    for s_i, member in enumerate(model.get("support") or []):
        sw = f"support[{s_i}] '{member.get('role', '')}'"
        if str(member.get("role", "")).lower() not in rates:
            errors.append(f"{sw}: role is not on the rate card")
        if not is_number(member.get("mandays_per_year"), strict=True):
            errors.append(f"{sw}: mandays_per_year must be above zero")
        check_cite(root, member.get("cite"), sw, errors)
    for r_i, run in enumerate(model.get("run_costs") or []):
        rw = f"run_costs[{r_i}] '{run.get('item', '')}'"
        if not run.get("item") or not is_number(run.get("annual_cost")):
            errors.append(f"{rw}: needs item and an annual_cost of zero or more")
        check_cite(root, run.get("cite"), rw, errors)
    for v_i, stream in enumerate(model.get("revenue") or []):
        vw = f"revenue[{v_i}] '{stream.get('stream', '')}'"
        amounts = stream.get("amounts")
        if not stream.get("stream"):
            errors.append(f"{vw}: stream (what earns it) is required")
        if not (isinstance(amounts, list) and amounts and all(is_number(x) for x in amounts)):
            errors.append(f"{vw}: amounts is a list of yearly revenue, zero or more each")
        elif not isinstance(horizon, int) or len(amounts) < horizon:
            errors.append(f"{vw}: give an amount for each of the {horizon} years of the horizon")
        elif len(amounts) > MAX_YEARS:
            errors.append(f"{vw}: at most {MAX_YEARS} yearly amounts")
        check_cite(root, stream.get("cite"), vw, errors)
    numbers = []
    for x_i, epic in enumerate(model.get("epics") or []):
        xw = f"epics[{x_i}] '{epic.get('epic', '')}'"
        if not isinstance(epic.get("number"), int) or isinstance(epic.get("number"), bool) or epic["number"] < 1:
            errors.append(f"{xw}: number is the epic's number (1, 2, ...), as in the story numbers 1.1, 1.2")
        elif epic["number"] in numbers:
            errors.append(f"{xw}: epic number {epic['number']} is used twice")
        numbers.append(epic.get("number"))
        if not epic.get("epic"):
            errors.append(f"{xw}: epic (its name) is required")
        if not (isinstance(epic.get("deliverables"), list) and all(isinstance(d, str) and d for d in epic["deliverables"])):
            errors.append(f"{xw}: deliverables is a list of the epic's stories or outputs")
        if not is_number(epic.get("days"), strict=True):
            errors.append(f"{xw}: days (working days) must be above zero")
        if epic.get("api_budget") is not None and not is_number(epic["api_budget"]):
            errors.append(f"{xw}: api_budget is zero or more, or left out")
        if epic.get("actual_team_cost") is not None and not is_number(epic["actual_team_cost"]):
            errors.append(f"{xw}: actual_team_cost (from timesheets) is zero or more, or left out")
        for key in ("actual_start", "actual_end"):
            if epic.get(key):
                try:
                    datetime.fromisoformat(str(epic[key])[:10])
                except ValueError:
                    errors.append(f"{xw}: {key} is a date, YYYY-MM-DD")
        check_cite(root, epic.get("cite"), xw, errors)
    if "sprint_days" in a and (not isinstance(a["sprint_days"], int) or a["sprint_days"] < 1):
        errors.append("assumptions.sprint_days: working days per sprint, a whole number of 1 or more")
    for p_i, sprint in enumerate(model.get("sprints") or []):
        pw = f"sprints[{p_i}]"
        if not isinstance(sprint.get("number"), int) or sprint["number"] < 1:
            errors.append(f"{pw}: number is the sprint's number (1, 2, ...)")
        if sprint.get("stories") is not None and not (isinstance(sprint["stories"], list) and all(isinstance(x, str) for x in sprint["stories"])):
            errors.append(f"{pw}: stories is a list of the stories planned for the sprint")
        for key in ("api_budget", "actual_team_cost"):
            if sprint.get(key) is not None and not is_number(sprint[key]):
                errors.append(f"{pw}: {key} is zero or more, or left out")
    phases = []
    for b_i, budget in enumerate(model.get("phase_budgets") or []):
        bw = f"phase_budgets[{b_i}] '{budget.get('phase', '')}'"
        if not isinstance(budget.get("phase"), str) or not budget["phase"].strip():
            errors.append(f"{bw}: phase is the planning phase as the token ledger names it (e.g. test design)")
        elif budget["phase"].strip().lower() in phases:
            errors.append(f"{bw}: phase '{budget['phase']}' is budgeted twice")
        else:
            phases.append(budget["phase"].strip().lower())
        if not is_number(budget.get("api_budget")):
            errors.append(f"{bw}: api_budget is zero or more")
        check_cite(root, budget.get("cite"), bw, errors)
    if a.get("start_date"):
        try:
            datetime.fromisoformat(str(a["start_date"]))
        except ValueError:
            errors.append("assumptions.start_date: a date, YYYY-MM-DD")
    for o_i, other in enumerate(model.get("other_costs") or []):
        ow = f"other_costs[{o_i}] '{other.get('item', '')}'"
        if other.get("epic") is not None and other["epic"] not in numbers:
            errors.append(f"{ow}: epic {other['epic']} is not one of the epics")
        if not other.get("item") or not is_number(other.get("cost")):
            errors.append(f"{ow}: needs item and a cost of zero or more")
        check_cite(root, other.get("cite"), ow, errors)
    if not isinstance(model.get("gaps", []), list):
        errors.append("gaps: a list of strings")
    return errors


# --------------------------------------------------------------------------- close-out (Claude usage, timeline)


def pricing_id(model: str) -> str:
    """A model id as the logs write it -> the pricing page's id: 'claude-haiku-4-5-20251001' -> 'claude-haiku-4-5'."""
    return re.sub(r"-\d{8}$", "", re.sub(r"\[.*?\]", "", model.strip().lower()))


def load_close_out(root: Path, model: dict) -> tuple[dict | None, list[str]]:
    """The Claude usage, Claude prices and project timeline the close-out sheets are built from."""
    spec = model.get("close_out")
    if spec is None:
        return None, []
    errors: list[str] = []
    ledger_dir = token_usage_folder(root)

    def path(key: str, default: Path) -> Path:
        value = spec.get(key)
        return (root / value) if value else default

    ledger, prices_file = path("ledger", ledger_dir / "claude_usage.csv"), path("prices", root / DATA_DIR / "claude-prices.json")
    timeline_file = path("timeline", ledger_dir / "project_timeline.csv")
    daily_file = path("daily", ledger_dir / "claude_usage_daily.csv")
    feedback_file = path("feedback", root / DATA_DIR / "ai-feedback.csv")
    usage, prices, timeline, daily, feedback = [], {}, [], [], []
    if not ledger.is_file():
        errors.append(f"close_out: no Claude usage ledger at {ledger}; refresh it with token_report.py")
    else:
        with ledger.open(newline="", encoding="utf-8") as fh:
            usage = [r for r in csv.DictReader(fh) if int(r.get("calls") or 0)]
    if not prices_file.is_file():
        errors.append(f"close_out: no Claude prices at {prices_file}; read them with claude_prices.py --save")
    else:
        prices = json.loads(prices_file.read_text(encoding="utf-8"))
    if timeline_file.is_file():
        with timeline_file.open(newline="", encoding="utf-8") as fh:
            timeline = list(csv.DictReader(fh))
    if daily_file.is_file():
        with daily_file.open(newline="", encoding="utf-8") as fh:
            daily = [r for r in csv.DictReader(fh) if int(r.get("calls") or 0)]
    if feedback_file.is_file():
        with feedback_file.open(newline="", encoding="utf-8") as fh:
            feedback = list(csv.DictReader(fh))
    ccy = model.get("currency")
    if ccy != "USD" and not is_number(spec.get("fx_rate"), strict=True):
        errors.append(f"close_out.fx_rate: Claude prices are in USD; give the rate to {ccy} (and fx_source)")
    known = prices.get("models", {})
    tokens = ("input_tokens", "cache_write_5m_tokens", "cache_write_1h_tokens", "cache_read_tokens", "output_tokens")
    unpriced = sorted({r["model"] for r in usage if pricing_id(r["model"]) not in known
                       and any(int(r.get(t) or 0) for t in tokens)})
    return {"usage": usage, "prices": prices, "timeline": timeline, "daily": daily, "feedback": feedback,
            "unpriced": unpriced,
            "fx_rate": spec.get("fx_rate", 1) if ccy != "USD" else 1, "fx_source": spec.get("fx_source", ""),
            "in_totals": spec.get("in_totals", True),
            "files": [p for p in (ledger, prices_file, timeline_file, daily_file, feedback_file) if p.is_file()]}, errors


def claude_cost(close: dict) -> tuple[float, int]:
    """(USD at API list prices, calls)."""
    known = close["prices"].get("models", {})
    usd, calls = 0.0, 0
    for r in close["usage"]:
        calls += int(r["calls"])
        price = known.get(pricing_id(r["model"]))
        if price:
            usd += sum(int(r[f"{k}_tokens"] or 0) * price[k] for k in
                       ("input", "cache_write_5m", "cache_write_1h", "cache_read", "output")) / 1_000_000
    return usd, calls


# --------------------------------------------------------------------------- figures


def settings(model: dict) -> dict:
    a = {**DEFAULTS, **model.get("assumptions", {})}
    a["contingency"] = a["contingency_pct"] / 100
    a["escalation"] = a["cost_escalation_pct"] / 100
    a["delivery_months"] = a["delivery_days"] / a["working_days_per_month"]
    h, e = a["horizon_years"], a["escalation"]
    a["horizon_factor"] = h if e == 0 else ((1 + e) ** h - 1) / e  # sum of (1+e)^(k-1) over the run years
    a["fx_rate"] = model.get("fx_rate", 1) if model.get("cloud_currency", model["currency"]) != model["currency"] else 1
    a["start"] = datetime.fromisoformat(a["start_date"]) if a.get("start_date") else None
    return a


def add_workdays(day, n: int):
    """Excel's WORKDAY: n working days (Monday to Friday) after day."""
    step = 1 if n >= 0 else -1
    while n:
        day += timedelta(days=step)
        if day.weekday() < 5:
            n -= step
    return day


def item_usage(item: dict, a: dict) -> float:
    usage = item.get("usage", "hours" if "hour" in str(item.get("unit", "")).lower() else None)
    return a["hours_per_month"] if usage == "hours" else usage


def compute(model: dict, card: dict, close: dict | None = None) -> dict:
    a = settings(model)
    h, e, cont = a["horizon_years"], a["escalation"], a["contingency"]
    rates = {str(r["role"]).lower(): r["day_rate"] for r in card.get("roles", [])}
    envs = []
    for env in model["environments"]:
        monthly = sum(i["unit_price"] * i["quantity"] * item_usage(i, a) * a["fx_rate"] for i in env["items"])
        d_months = a["delivery_months"] if env.get("delivery_months") is None else env["delivery_months"]
        per_year = monthly * env.get("run_months_per_year", 12)
        envs.append({"name": env["name"], "monthly": monthly, "daily": monthly * 12 / 365, "delivery": monthly * d_months,
                     "run_per_year": per_year, "run": per_year * a["horizon_factor"]})
    mandays = people = 0.0
    for m in model.get("team") or []:
        days = a["delivery_days"] if m.get("days") is None else m["days"]
        md = m["employees"] * days * m.get("allocation", 1)
        mandays += md
        people += md * rates[m["role"].lower()]
    support_md = sum(m["mandays_per_year"] for m in model.get("support") or [])
    support = sum(m["mandays_per_year"] * rates[m["role"].lower()] for m in model.get("support") or [])
    run_costs = sum(r["annual_cost"] for r in model.get("run_costs") or [])
    other = sum(o["cost"] for o in model.get("other_costs") or [])
    claude = None
    if close:
        usd, calls = claude_cost(close)
        claude = {"calls": calls, "usd": round(usd, 2), "cost": round(usd * close["fx_rate"], 2),
                  "in_totals": bool(close["in_totals"]), "unpriced_models": close["unpriced"]}
    delivery = sum(x["delivery"] for x in envs) + people + other + (claude["cost"] if claude and claude["in_totals"] else 0)
    run_per_year = sum(x["run_per_year"] for x in envs) + support + run_costs
    run = run_per_year * a["horizon_factor"]

    # Multi-year view: the build, then years 1..5 (zero past the horizon), contingency on every cost
    revenue = [0.0] * MAX_YEARS
    for stream in model.get("revenue") or []:
        for k, amount in enumerate(stream["amounts"][:MAX_YEARS]):
            revenue[k] += amount
    rev = [0.0] + [revenue[k] if k < h else 0.0 for k in range(MAX_YEARS)]
    costs = [delivery * (1 + cont)] + [run_per_year * (1 + e) ** k * (1 + cont) if k < h else 0.0 for k in range(MAX_YEARS)]
    net = [r_ - c for r_, c in zip(rev, costs)]
    cumulative, running = [], 0.0
    for n in net:
        running += n
        cumulative.append(running)
    epic_list = model.get("epics") or []
    epic_days = sum(x["days"] for x in epic_list)
    env_delivery = sum(x["delivery"] for x in envs)
    story_cost: dict[str, float] = {}
    if close:
        known = close["prices"].get("models", {})
        for u in close["usage"]:
            price = known.get(pricing_id(u["model"]))
            if price and u.get("story_no"):
                story_cost[u["story_no"]] = story_cost.get(u["story_no"], 0) + sum(
                    int(u[f"{k_}_tokens"] or 0) * price[k_] for k_ in
                    ("input", "cache_write_5m", "cache_write_1h", "cache_read", "output")) / 1e6 * close["fx_rate"]
    timeline = {t["name"]: t for t in (close or {}).get("timeline", []) if t.get("category") == "epic"}
    epics_out, day = [], a["start"]
    for x in epic_list:
        share = x["days"] / epic_days
        team_x, infra_x = people * share, env_delivery * share
        other_x = sum(o["cost"] for o in model.get("other_costs") or [] if o.get("epic") == x["number"])
        actual_api = sum(v for k_, v in story_cost.items() if k_.startswith(f"{x['number']}."))
        start = end = None
        if day:
            start = day if not epics_out else add_workdays(day, 1)
            end = add_workdays(start, int(x["days"]) - 1)
            day = end
        forecast = team_x + infra_x + other_x + (x.get("api_budget") or 0)
        t = timeline.get(f"epic-{x['number']}", {})
        a_start = x.get("actual_start") or t.get("started")
        a_end = x.get("actual_end") or (t.get("completed") if t.get("status", "done") == "done" else None)
        out_x = {"number": x["number"], "epic": x["epic"], "days": x["days"], "weeks": x["days"] / 5,
                 "start": start.date().isoformat() if start else None, "end": end.date().isoformat() if end else None,
                 "share": share, "mandays": mandays * share, "team_cost": team_x, "infra_cost": infra_x,
                 "other_costs": other_x, "api_budget": x.get("api_budget"), "forecast_total": forecast,
                 "actual_start": a_start[:10] if a_start else None, "actual_end": a_end[:10] if a_end else None,
                 "actual_days": None, "days_variance": None, "days_variance_pct": None,
                 "api_actual": actual_api if close else None, "actual_team_cost": None, "actual_total": None,
                 "cost_variance": None}
        if a_start and a_end:
            s_, e_ = datetime.fromisoformat(a_start[:10]), datetime.fromisoformat(a_end[:10])
            days_a = sum(1 for k_ in range((e_ - s_).days + 1) if (s_ + timedelta(days=k_)).weekday() < 5)
            ratio = days_a / x["days"]
            team_a = x["actual_team_cost"] if x.get("actual_team_cost") is not None else team_x * ratio
            total_a = team_a + infra_x * ratio + other_x + actual_api
            out_x.update(actual_days=days_a, days_variance=days_a - x["days"], days_variance_pct=ratio - 1,
                         actual_team_cost=team_a, actual_total=total_a, cost_variance=total_a - forecast)
        epics_out.append(out_x)
    # Sprints: fixed timeboxes of sprint_days working days from the start date; actuals up to the last actual activity
    sprints_out = []
    planned = {x["number"]: x for x in model.get("sprints") or []}
    sd, dd = int(a["sprint_days"]), a["delivery_days"]
    known = (close or {}).get("prices", {}).get("models", {})
    daily_cost = []
    for u in (close or {}).get("daily") or []:
        price = known.get(pricing_id(u["model"]))
        cost = sum(int(u.get(f"{k_}_tokens") or 0) * price[k_] for k_ in
                   ("input", "cache_write_5m", "cache_write_1h", "cache_read", "output")) / 1e6 if price else 0
        daily_cost.append((u["day"][:10], cost * close["fx_rate"]))
    done_stories = [(t["completed"][:10], t["name"]) for t in (close or {}).get("timeline") or []
                    if t.get("category") == "story" and t.get("completed")]
    last_actual = max([d for d, _ in done_stories] + [d for d, _ in daily_cost], default=None)
    burn, infra_burn = people / dd, env_delivery / dd
    k = 0
    while True:
        k += 1
        planned_days = min(sd, max(0, dd - (k - 1) * sd))
        start = add_workdays(a["start"], (k - 1) * sd) if a["start"] else None
        reached = bool(start and last_actual and start.date().isoformat() <= last_actual)
        if planned_days <= 0 and not reached:
            break
        end = add_workdays(start, sd - 1) if start else None
        plan = planned.get(k, {})
        forecast = planned_days * (burn + infra_burn) + (plan.get("api_budget") or 0)
        row = {"number": k, "start": start.date().isoformat() if start else None, "end": end.date().isoformat() if end else None,
               "days": planned_days, "planned_stories": plan.get("stories") or [], "team_cost": planned_days * burn,
               "infra_cost": planned_days * infra_burn, "api_budget": plan.get("api_budget"), "forecast_total": forecast,
               "days_worked": None, "stories_done": None, "stories_done_list": [], "api_actual": None,
               "actual_team_cost": None, "actual_total": None, "cost_variance": None}
        if reached:
            s_, e_ = row["start"], row["end"]
            stop = min(end.date(), datetime.fromisoformat(last_actual).date())
            worked = sum(1 for d_ in range((stop - start.date()).days + 1) if (start.date() + timedelta(days=d_)).weekday() < 5)
            api = sum(c for d, c in daily_cost if s_ <= d <= e_)
            team_a = plan["actual_team_cost"] if plan.get("actual_team_cost") is not None else worked * burn
            total_a = team_a + worked * infra_burn + api
            row["stories_done_list"] = [n for d, n in sorted(done_stories) if s_ <= d <= e_]
            row.update(days_worked=worked, stories_done=len(row["stories_done_list"]) if done_stories else None,
                       api_actual=api if daily_cost else None, actual_team_cost=team_a, actual_total=total_a,
                       cost_variance=total_a - forecast)
        sprints_out.append(row)

    # Phase budgets: the Claude API budget per planning phase against the ledger's actual for that work item
    item_cost: dict[str, float] = {}
    for u in (close or {}).get("usage") or []:
        price = known.get(pricing_id(u["model"]))
        if price:
            item_cost[u["work_item"].lower()] = item_cost.get(u["work_item"].lower(), 0) + sum(
                int(u.get(f"{k_}_tokens") or 0) * price[k_] for k_ in
                ("input", "cache_write_5m", "cache_write_1h", "cache_read", "output")) / 1e6 * close["fx_rate"]
    phases_out = []
    for b in model.get("phase_budgets") or []:
        actual = item_cost.get(b["phase"].strip().lower(), 0.0) if close else None
        phases_out.append({"phase": b["phase"], "api_budget": b["api_budget"], "api_actual": actual,
                           "variance": None if actual is None else actual - b["api_budget"],
                           "variance_pct": None if actual is None or not b["api_budget"] else actual / b["api_budget"] - 1})

    periods = ["Build"] + [f"Year {k + 1}" for k in range(MAX_YEARS)]
    payback = next((p_ for p_, c in zip(periods, cumulative) if c >= 0), "beyond the horizon")
    r2 = lambda x: round(x, 2)  # noqa: E731
    total_costs = sum(costs)
    return {
        "currency": model["currency"],
        "horizon_years": h,
        "delivery_months": r2(a["delivery_months"]),
        "environments": [{k: (r2(v) if isinstance(v, float | int) else v) for k, v in x.items()} for x in envs],
        "team": {"employees": sum(m["employees"] for m in model.get("team") or []), "mandays": r2(mandays),
                 "cost": r2(people)},
        "support": {"mandays_per_year": r2(support_md), "cost_per_year": r2(support)},
        "run_costs_per_year": r2(run_costs),
        "infra_opex": {k: r2(v) for k, v in infra_rows(model)[-1].items() if k in ("daily", "weekly", "monthly", "annual")},
        "other_costs": r2(other),
        **({"claude": claude} if claude else {}),
        "subtotal": {"delivery": r2(delivery), "run_per_year": r2(run_per_year), "run": r2(run), "total": r2(delivery + run)},
        "contingency": r2((delivery + run) * cont),
        "grand_total": {"delivery": r2(delivery * (1 + cont)), "run": r2(run * (1 + cont)),
                        "total": r2((delivery + run) * (1 + cont))},
        "multi_year": {"periods": periods, "revenue": [r2(x) for x in rev], "costs": [r2(x) for x in costs],
                       "net": [r2(x) for x in net], "cumulative": [r2(x) for x in cumulative]},
        "revenue_total": r2(sum(rev)),
        "net_total": r2(sum(net)),
        "payback": payback,
        "roi": round(sum(net) / total_costs, 4) if total_costs else None,
        "epics": [{k_: (round(v, 4) if k_ in ("share", "days_variance_pct") else r2(v)) if isinstance(v, float) else v
                   for k_, v in x.items()} for x in epics_out],
        "sprints": [{k_: r2(v) if isinstance(v, float) else v for k_, v in x.items()} for x in sprints_out],
        "phase_budgets": [{k_: (round(v, 4) if k_ == "variance_pct" else r2(v)) if isinstance(v, float) else v
                           for k_, v in x.items()} for x in phases_out],
        "gaps": len(model.get("gaps") or []),
    }


# --------------------------------------------------------------------------- stamp and drift


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def cited_files(model: dict) -> list[str]:
    cites = [i.get("cite", "") for e in model.get("environments", []) for i in e.get("items", [])]
    cites += [m.get("cite", "") for m in model.get("team") or []]
    for key in ("other_costs", "run_costs", "support", "revenue", "phase_budgets"):
        cites += [o.get("cite", "") for o in model.get(key) or []]
    cites += [x.get("cite", "") for x in model.get("epics") or []]
    files = {split_cite(c)[1] for c in cites if isinstance(c, str)}
    return sorted(f for f in files if f)


def make_stamp(root: Path, model: dict, card_path: Path | None, extra: list[Path] = ()) -> dict:
    files = {f: sha(root / f) for f in cited_files(model)}
    for path in extra:
        files[path.resolve().relative_to(root.resolve()).as_posix()] = sha(path)
    if card_path and card_path.is_file():
        rel = card_path.resolve().relative_to(root.resolve()).as_posix() if card_path.resolve().is_relative_to(
            root.resolve()) else str(card_path)
        files[rel] = sha(card_path)
    return {"built": datetime.now().isoformat(timespec="seconds"), "files": files}


def drift(root: Path, model: dict) -> dict:
    stamp = model.get("stamp")
    if not stamp:
        return {"status": "never-built", "changed": [], "missing": []}
    changed, missing = [], []
    for rel, digest in stamp.get("files", {}).items():
        path = Path(rel) if Path(rel).is_absolute() else root / rel
        if not path.is_file():
            missing.append(rel)
        elif sha(path) != digest:
            changed.append(rel)
    return {"status": "drift" if changed or missing else "current", "built": stamp.get("built"),
            "changed": changed, "missing": missing}


# --------------------------------------------------------------------------- workbook


def write_workbook(model: dict, card: dict, stamp: dict, out: Path, close: dict | None = None,
                   sprint_figures: list | None = None) -> list[str]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table, TableStyleInfo

    a = settings(model)
    ccy = model["currency"]
    cloud_ccy = model.get("cloud_currency", ccy)
    region = model.get("pricing_region", "")
    price_date = model.get("price_date", "")

    MONEY, QTY, PCT, MONTHS = "#,##0.00", "#,##0.##", "0%", "0.0"
    BLUE, GREEN = Font(color="0000FF"), Font(color="008000")
    BOLD = Font(bold=True)
    HEAD = Font(bold=True, color="FFFFFF")
    HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
    SECTION = Font(bold=True, size=12, color="1F4E78")
    TITLE = Font(bold=True, size=14)
    TOP = Border(top=Side(style="thin"))
    WRAP = Alignment(wrap_text=True, vertical="top")

    def q(sheet: str) -> str:
        return "'" + sheet.replace("'", "''") + "'"

    def header(ws, row: int, labels: list[str]) -> None:
        for c, label in enumerate(labels, 1):
            cell = ws.cell(row=row, column=c, value=label)
            cell.font, cell.fill, cell.alignment = HEAD, HEAD_FILL, Alignment(wrap_text=True, vertical="center")

    def put(ws, ref: str, value, fmt: str | None = None, font=None):
        cell = ws[ref]
        cell.value = value
        if fmt:
            cell.number_format = fmt
        if font is not None:
            cell.font = font
        return cell

    def widths(ws, values: list[int]) -> None:
        for i, w in enumerate(values, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    tables: set[str] = set()

    def add_table(ws, name: str, first_col: str, header_row: int, last_col: str, last_row: int) -> None:
        """Name a data range as an Excel table so Power BI and Excel Online pick it up by name."""
        name = re.sub(r"[^A-Za-z0-9_]", "", name)
        base, n = name, 2
        while name.lower() in tables:
            name, n = f"{base}{n}", n + 1
        tables.add(name.lower())
        t = Table(displayName=name, ref=f"{first_col}{header_row}:{last_col}{max(last_row, header_row + 1)}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(t)

    def total_row(ws, row: int, last_col: int) -> None:
        for c in range(1, last_col + 1):
            ws.cell(row=row, column=c).border = TOP
            ws.cell(row=row, column=c).font = BOLD

    wb = Workbook()
    ov = wb.active
    ov.title = "Overall"

    YEARS = [f"Year {k}" for k in range(1, MAX_YEARS + 1)]
    YEAR_COLS = "DEFGH"  # Multi-Year and Revenue: Year 1..5

    # --- Environments (built first so Overall can point at their totals)
    env_totals: list[tuple[str, int]] = []
    for env in model["environments"]:
        ws = wb.create_sheet(env["name"])
        put(ws, "A1", f"{env['name']} environment", font=TITLE)
        fx_note = f", converted to {ccy} at the Overall exchange rate" if cloud_ccy != ccy else ""
        put(ws, "A2", f"Azure retail list prices (pay-as-you-go) in {region}, {cloud_ccy}, fetched {price_date}"
                      f"{fx_note}. Daily cost is the monthly cost spread over the year (x 12 / 365). Blue cells are inputs.")
        header(ws, 4, ["#", "Component", "Azure service", "SKU / tier", "Unit", f"Unit price ({cloud_ccy})",
                       "Quantity", "Usage per month", f"Monthly cost ({ccy})", f"Daily cost ({ccy})",
                       f"Annual cost ({ccy})", "Cited from", "Price source", "Note", "Environment"])
        row = 5
        for n, item in enumerate(env["items"], 1):
            usage = item.get("usage", "hours" if "hour" in str(item.get("unit", "")).lower() else None)
            ws.cell(row=row, column=1, value=n)
            ws.cell(row=row, column=2, value=item["component"])
            ws.cell(row=row, column=3, value=item["service"])
            ws.cell(row=row, column=4, value=item.get("sku", ""))
            ws.cell(row=row, column=5, value=item["unit"])
            put(ws, f"F{row}", item["unit_price"], "#,##0.0000", BLUE)
            put(ws, f"G{row}", item["quantity"], QTY, BLUE)
            if usage == "hours":
                put(ws, f"H{row}", "=HoursPerMonth", QTY)
            else:
                put(ws, f"H{row}", usage, QTY, BLUE)
            put(ws, f"I{row}", f"=F{row}*G{row}*H{row}*FxRate", MONEY)
            put(ws, f"J{row}", f"=I{row}*12/365", MONEY)
            put(ws, f"K{row}", f"=I{row}*12", MONEY)
            ws.cell(row=row, column=12, value=item["cite"])
            ws.cell(row=row, column=13, value=item["price_source"])
            ws.cell(row=row, column=14, value=item.get("note", ""))
            ws.cell(row=row, column=15, value=env["name"])
            for c in (12, 13, 14):
                ws.cell(row=row, column=c).alignment = WRAP
            row += 1
        for _ in range(SPARE_ROWS):
            for col in "FGH":
                ws[f"{col}{row}"].font = BLUE
            put(ws, f"I{row}", f'=IF(F{row}="","",F{row}*G{row}*H{row}*FxRate)', MONEY)
            put(ws, f"J{row}", f'=IF(I{row}="","",I{row}*12/365)', MONEY)
            put(ws, f"K{row}", f'=IF(I{row}="","",I{row}*12)', MONEY)
            ws.cell(row=row, column=15, value=env["name"])
            row += 1
        add_table(ws, f"tblEnv{env['name']}", "A", 4, "O", row - 1)
        put(ws, f"A{row}", "Total")
        for col in "IJK":
            put(ws, f"{col}{row}", f"=SUM({col}5:{col}{row - 1})", MONEY)
        total_row(ws, row, 14)
        env_totals.append((env["name"], row))
        widths(ws, [5, 26, 26, 24, 12, 14, 10, 12, 16, 14, 16, 40, 48, 30, 14])
        ws.freeze_panes = "A5"

    # --- Rate Card
    rc = wb.create_sheet("Rate Card")
    put(rc, "A1", "Employee rate card", font=TITLE)
    put(rc, "A2", f"Day rates in {ccy}. Source: {card.get('source', 'not stated')}; updated "
                  f"{card.get('updated', 'not stated')}. Blue cells are inputs; the team and support tables look roles up here.")
    header(rc, 4, ["Role", f"Day rate ({ccy})", "Note"])
    r = 5
    for role in card.get("roles", []):
        rc.cell(row=r, column=1, value=role["role"])
        put(rc, f"B{r}", role["day_rate"], MONEY, BLUE)
        rc.cell(row=r, column=3, value=role.get("note", ""))
        r += 1
    add_table(rc, "tblRateCard", "A", 4, "C", r - 1)
    rate_last = max(r + 20, 50)  # room for roles added by hand
    widths(rc, [30, 16, 50])
    rc.freeze_panes = "A5"

    def simple_sheet(title: str, heading: str, note: str, amount_label: str, rows: list[tuple], table: str,
                     epics: list | None = None) -> int:
        """# | Item | Amount | Cited from | Note [| Epic], with spare rows and a total; returns the total's row."""
        ws = wb.create_sheet(title)
        put(ws, "A1", heading, font=TITLE)
        put(ws, "A2", note)
        header(ws, 4, ["#", "Item", amount_label, "Cited from", "Note"] + (["Epic"] if epics is not None else []))
        last_col = "F" if epics is not None else "E"
        r = 5
        for n, (item, amount, cite, extra) in enumerate(rows, 1):
            ws.cell(row=r, column=1, value=n)
            ws.cell(row=r, column=2, value=item)
            put(ws, f"C{r}", amount, MONEY, BLUE)
            ws.cell(row=r, column=4, value=cite)
            ws.cell(row=r, column=5, value=extra)
            if epics is not None and epics[n - 1] is not None:
                put(ws, f"F{r}", epics[n - 1], "0", BLUE)
            r += 1
        for _ in range(SPARE_ROWS):
            ws[f"C{r}"].font = BLUE
            if epics is not None:
                ws[f"F{r}"].font = BLUE
            r += 1
        add_table(ws, table, "A", 4, last_col, r - 1)
        put(ws, f"A{r}", "Total")
        put(ws, f"C{r}", f"=SUM(C5:C{r - 1})", MONEY)
        total_row(ws, r, 6 if epics is not None else 5)
        widths(ws, [5, 36, 18, 44, 30, 8])
        return r

    other_total = simple_sheet(
        "Other Costs", "Other one-off costs",
        f"Licences, tests, training and other one-off costs of the build, outside the environments and the team, in {ccy}.",
        f"Amount ({ccy})", [(o["item"], o["cost"], o["cite"], o.get("note", "")) for o in model.get("other_costs") or []],
        "tblOtherCosts", epics=[o.get("epic") for o in model.get("other_costs") or []])
    run_costs_total = simple_sheet(
        "Run Costs", "Recurring run costs (OPEX)",
        f"Yearly costs of running the solution outside the cloud environments and the support team: support contracts, "
        f"licences, subscriptions. In {ccy} per year; the Multi-Year sheet escalates them by the yearly increase.",
        f"Annual cost ({ccy})", [(x["item"], x["annual_cost"], x["cite"], x.get("note", "")) for x in model.get("run_costs") or []],
        "tblRunCosts")

    # --- Revenue
    rv = wb.create_sheet("Revenue")
    put(rv, "A1", "Revenue", font=TITLE)
    put(rv, "A2", f"Revenue the solution earns, per year after go-live, in {ccy}, from the owner or the business case; each "
                  "stream cites its source. Years past the horizon are ignored. Blue cells are inputs.")
    header(rv, 4, ["#", "Stream", "Basis", *YEARS, "Cited from", "Note"])
    r = 5
    for n, stream in enumerate(model.get("revenue") or [], 1):
        rv.cell(row=r, column=1, value=n)
        rv.cell(row=r, column=2, value=stream["stream"])
        rv.cell(row=r, column=3, value=stream.get("basis", ""))
        for k, col in enumerate(YEAR_COLS):
            amount = stream["amounts"][k] if k < len(stream["amounts"]) else 0
            put(rv, f"{col}{r}", amount, MONEY, BLUE)
        rv.cell(row=r, column=9, value=stream["cite"])
        rv.cell(row=r, column=10, value=stream.get("note", ""))
        r += 1
    for _ in range(SPARE_ROWS):
        for col in YEAR_COLS:
            rv[f"{col}{r}"].font = BLUE
        r += 1
    add_table(rv, "tblRevenue", "A", 4, "J", r - 1)
    revenue_first, revenue_last = 5, r - 1
    put(rv, f"A{r}", "Total")
    for col in YEAR_COLS:
        put(rv, f"{col}{r}", f"=SUM({col}5:{col}{r - 1})", MONEY)
    total_row(rv, r, 10)
    widths(rv, [5, 30, 28, 14, 14, 14, 14, 14, 40, 30])

    # --- Overall: inputs
    put(ov, "A1", model.get("title") or f"{model['name']} cost estimate", font=TITLE)
    put(ov, "A2", f"All costs in {ccy}. Cloud: Azure retail list prices in {region} fetched {price_date}, "
                  "no discounts or reservations. Blue cells are inputs; black cells are formulas; green cells link to other sheets.")
    put(ov, "A3", "Inputs", font=SECTION)
    header(ov, 4, ["Input", "Value"])
    inputs = [
        ("DeliveryDays", "Days of delivery (working days)", a["delivery_days"], QTY, True),
        ("WorkingDaysPerMonth", "Working days per month", a["working_days_per_month"], QTY, True),
        ("DeliveryMonths", "Months of delivery", "=DeliveryDays/WorkingDaysPerMonth", MONTHS, False),
        ("HoursPerMonth", "Hours per month (always-on services)", a["hours_per_month"], QTY, True),
        ("ProjectStart", "Delivery start date (optional; dates the epics and sprints)", a.get("start"), "yyyy-mm-dd", True),
        ("SprintDays", "Working days per sprint", a["sprint_days"], "0", True),
        ("HorizonYears", "Years of run after go-live (3 to 5)", a["horizon_years"], "0", True),
        ("CostEscalation", "Yearly increase in run costs", a["escalation"], "0.0%", True),
        ("HorizonFactor", "Run years, with the yearly increase",
         "=IF(CostEscalation=0,HorizonYears,((1+CostEscalation)^HorizonYears-1)/CostEscalation)", "0.00", False),
        ("Contingency", "Contingency", a["contingency"], PCT, True),
        ("FxRate", f"Exchange rate (1 {cloud_ccy} = x {ccy})", a["fx_rate"], "0.0000", True),
    ]
    if close:
        inputs.append(("ClaudeFxRate", f"Exchange rate for Claude prices (1 USD = x {ccy})", close["fx_rate"], "0.0000", True))
    input_row = {}
    for offset, (name, label, value, fmt, is_input) in enumerate(inputs):
        r = 5 + offset
        input_row[name] = r
        ov.cell(row=r, column=1, value=label)
        put(ov, f"B{r}", value, fmt, BLUE if is_input else None)
        wb.defined_names[name] = DefinedName(name, attr_text=f"Overall!$B${r}")
    add_table(ov, "tblInputs", "A", 4, "B", 4 + len(inputs))
    horizon_check = DataValidation(type="whole", operator="between", formula1="3", formula2=str(MAX_YEARS),
                                   showErrorMessage=True, error=f"3 to {MAX_YEARS} years")
    ov.add_data_validation(horizon_check)
    horizon_check.add(f"B{input_row['HorizonYears']}")
    if cloud_ccy != ccy:
        ov.cell(row=input_row["FxRate"], column=3, value=model.get("fx_source", ""))
    if close and ccy != "USD":
        ov.cell(row=input_row["ClaudeFxRate"], column=3, value=close["fx_source"])

    # --- Overall: delivery team
    r = 5 + len(inputs) + 1
    put(ov, f"A{r}", "Delivery team (build)", font=SECTION)
    r += 1
    header(ov, r, ["Role", f"Day rate ({ccy})", "Employees", "Days on delivery", "Allocation", "Mandays",
                   f"Cost ({ccy})", "Cited from"])
    first = r + 1
    rate_range = f"{q('Rate Card')}!$A$5:$B${rate_last}"
    for m in model.get("team") or []:
        r += 1
        ov.cell(row=r, column=1, value=m["role"])
        put(ov, f"B{r}", f"=VLOOKUP(A{r},{rate_range},2,FALSE)", MONEY, GREEN)
        put(ov, f"C{r}", m["employees"], QTY, BLUE)
        if m.get("days") is None:
            put(ov, f"D{r}", "=DeliveryDays", QTY)
        else:
            put(ov, f"D{r}", m["days"], QTY, BLUE)
        put(ov, f"E{r}", m.get("allocation", 1), PCT, BLUE)
        put(ov, f"F{r}", f"=C{r}*D{r}*E{r}", QTY)
        put(ov, f"G{r}", f"=F{r}*B{r}", MONEY)
        ov.cell(row=r, column=8, value=m["cite"])
    for _ in range(SPARE_ROWS):
        r += 1
        put(ov, f"B{r}", f'=IF(A{r}="","",VLOOKUP(A{r},{rate_range},2,FALSE))', MONEY, GREEN)
        for col in "CDE":
            ov[f"{col}{r}"].font = BLUE
        ov[f"E{r}"].number_format = PCT
        put(ov, f"F{r}", f'=IF(A{r}="","",C{r}*D{r}*E{r})', QTY)
        put(ov, f"G{r}", f'=IF(A{r}="","",F{r}*B{r})', MONEY)
    add_table(ov, "tblTeam", "A", first - 1, "H", r)
    r += 1
    team_total = r
    put(ov, f"A{r}", "Total")
    put(ov, f"C{r}", f"=SUM(C{first}:C{r - 1})", QTY)
    put(ov, f"F{r}", f"=SUM(F{first}:F{r - 1})", QTY)
    put(ov, f"G{r}", f"=SUM(G{first}:G{r - 1})", MONEY)
    total_row(ov, r, 8)

    # --- Overall: run support (OPEX in mandays)
    r += 2
    put(ov, f"A{r}", "Run support (OPEX): IT support and maintenance after go-live", font=SECTION)
    r += 1
    header(ov, r, ["Role", f"Day rate ({ccy})", "Mandays per year", f"Cost per year ({ccy})", "Cited from"])
    first = r + 1
    for m in model.get("support") or []:
        r += 1
        ov.cell(row=r, column=1, value=m["role"])
        put(ov, f"B{r}", f"=VLOOKUP(A{r},{rate_range},2,FALSE)", MONEY, GREEN)
        put(ov, f"C{r}", m["mandays_per_year"], QTY, BLUE)
        put(ov, f"D{r}", f"=C{r}*B{r}", MONEY)
        ov.cell(row=r, column=5, value=m["cite"])
    for _ in range(SPARE_ROWS):
        r += 1
        put(ov, f"B{r}", f'=IF(A{r}="","",VLOOKUP(A{r},{rate_range},2,FALSE))', MONEY, GREEN)
        ov[f"C{r}"].font = BLUE
        put(ov, f"D{r}", f'=IF(A{r}="","",C{r}*B{r})', MONEY)
    add_table(ov, "tblSupport", "A", first - 1, "E", r)
    r += 1
    support_total = r
    put(ov, f"A{r}", "Total")
    put(ov, f"C{r}", f"=SUM(C{first}:C{r - 1})", QTY)
    put(ov, f"D{r}", f"=SUM(D{first}:D{r - 1})", MONEY)
    total_row(ov, r, 5)

    # --- Overall: cost summary
    r += 2
    put(ov, f"A{r}", "Cost summary", font=SECTION)
    r += 1
    header(ov, r, ["Item", f"Monthly ({ccy})", "Months in delivery", f"Delivery cost ({ccy})",
                   "Run months per year", f"Run cost per year ({ccy})", f"Run cost over the horizon ({ccy})",
                   f"Total ({ccy})", "From"])
    first = r + 1
    run_rows: list[tuple[str, int]] = []  # (label, row) of every recurring cost, for the Multi-Year sheet
    for env, (name, env_total) in zip(model["environments"], env_totals):
        r += 1
        ov.cell(row=r, column=1, value=f"{name} environment")
        put(ov, f"B{r}", f"={q(name)}!$I${env_total}", MONEY, GREEN)
        if env.get("delivery_months") is None:
            put(ov, f"C{r}", "=DeliveryMonths", MONTHS)
        else:
            put(ov, f"C{r}", env["delivery_months"], MONTHS, BLUE)
        put(ov, f"D{r}", f"=B{r}*C{r}", MONEY)
        put(ov, f"E{r}", env.get("run_months_per_year", 12), MONTHS, BLUE)
        put(ov, f"F{r}", f"=B{r}*E{r}", MONEY)
        put(ov, f"G{r}", f"=F{r}*HorizonFactor", MONEY)
        put(ov, f"H{r}", f"=D{r}+G{r}", MONEY)
        ov.cell(row=r, column=9, value=f"{name} sheet")
        run_rows.append((f"{name} environment (run)", r))
    r += 1
    ov.cell(row=r, column=1, value="Delivery team")
    put(ov, f"D{r}", f"=G{team_total}", MONEY, GREEN)
    put(ov, f"H{r}", f"=D{r}+G{r}", MONEY)
    ov.cell(row=r, column=9, value="Delivery team table above")
    r += 1
    ov.cell(row=r, column=1, value="Run support (OPEX)")
    put(ov, f"F{r}", f"=D{support_total}", MONEY, GREEN)
    put(ov, f"G{r}", f"=F{r}*HorizonFactor", MONEY)
    put(ov, f"H{r}", f"=D{r}+G{r}", MONEY)
    ov.cell(row=r, column=9, value="Run support table above")
    run_rows.append(("Run support (OPEX)", r))
    r += 1
    ov.cell(row=r, column=1, value="Recurring run costs")
    put(ov, f"F{r}", f"={q('Run Costs')}!$C${run_costs_total}", MONEY, GREEN)
    put(ov, f"G{r}", f"=F{r}*HorizonFactor", MONEY)
    put(ov, f"H{r}", f"=D{r}+G{r}", MONEY)
    ov.cell(row=r, column=9, value="Run Costs sheet")
    run_rows.append(("Recurring run costs", r))
    r += 1
    ov.cell(row=r, column=1, value="Other one-off costs")
    put(ov, f"D{r}", f"={q('Other Costs')}!$C${other_total}", MONEY, GREEN)
    put(ov, f"H{r}", f"=D{r}+G{r}", MONEY)
    ov.cell(row=r, column=9, value="Other Costs sheet")
    claude_ref = None
    if close:
        claude_ref = "=" + q("Claude Usage") + "!$I$" + str(5 + len(claude_models(close)))
        if close["in_totals"]:
            r += 1
            ov.cell(row=r, column=1, value="Claude API usage (actual, API list prices)")
            put(ov, f"D{r}", claude_ref, MONEY, GREEN)
            put(ov, f"H{r}", f"=D{r}+G{r}", MONEY)
            ov.cell(row=r, column=9, value="Claude Usage sheet")
    add_table(ov, "tblCostSummary", "A", first - 1, "I", r)
    r += 1
    sub = r
    put(ov, f"A{r}", "Subtotal")
    for col in "DFGH":
        put(ov, f"{col}{r}", f"=SUM({col}{first}:{col}{r - 1})", MONEY)
    total_row(ov, r, 9)
    r += 1
    ov.cell(row=r, column=1, value="Contingency")
    for col in "DFG":
        put(ov, f"{col}{r}", f"={col}{sub}*Contingency", MONEY)
    put(ov, f"H{r}", f"=D{r}+G{r}", MONEY)
    r += 1
    grand = r
    put(ov, f"A{r}", "Grand total")
    for col in "DFGH":
        put(ov, f"{col}{r}", f"={col}{sub}+{col}{r - 1}", MONEY)
    total_row(ov, r, 9)
    for col in "ADFGH":
        ov[f"{col}{r}"].font = Font(bold=True, size=12)
    if close and not close["in_totals"]:
        r += 2
        ov.cell(row=r, column=1, value="Memo: Claude API usage at API list prices (not in the totals: Claude Code was on a subscription)")
        put(ov, f"D{r}", claude_ref, MONEY, GREEN)
    widths(ov, [40, 16, 14, 18, 14, 18, 20, 18, 30])

    # --- Multi-Year: the build, then years 1..5 (zero past the horizon)
    my = wb.create_sheet("Multi-Year")
    put(my, "A1", "Multi-year view", font=TITLE)
    put(my, "A2", f"In {ccy}. Build is the delivery; each year after go-live carries the run costs, raised by the yearly increase, "
                  "and the revenue given for it. Years past the horizon (Overall input) show zero. Contingency applies to every cost.")
    header(my, 4, ["Category", "Item", "Build", *YEARS, "Total"])
    r = 4
    ovq = q("Overall")

    def year_cells(row: int, formula) -> None:
        for k, col in enumerate(YEAR_COLS, 1):
            put(my, f"{col}{row}", f"=IF({k}<=HorizonYears,{formula(k)},0)", MONEY)
        put(my, f"I{row}", f"=SUM(C{row}:H{row})", MONEY)

    for vr in range(revenue_first, revenue_last + 1):
        r += 1
        my.cell(row=r, column=1, value="Revenue")
        put(my, f"B{r}", f'=IF({q("Revenue")}!$B${vr}="","(spare revenue row)",{q("Revenue")}!$B${vr})', None, GREEN)
        put(my, f"C{r}", 0, MONEY)
        year_cells(r, lambda k, vr=vr: f"{q('Revenue')}!{YEAR_COLS[k - 1]}${vr}")
    r += 1
    my.cell(row=r, column=1, value="Cost")
    my.cell(row=r, column=2, value="Build: delivery team, environments during delivery, one-off costs"
                                   + (", Claude usage" if close and close["in_totals"] else ""))
    put(my, f"C{r}", f"={ovq}!$D${sub}", MONEY, GREEN)
    for col in YEAR_COLS:
        put(my, f"{col}{r}", 0, MONEY)
    put(my, f"I{r}", f"=SUM(C{r}:H{r})", MONEY)
    for label, orow in run_rows:
        r += 1
        my.cell(row=r, column=1, value="Cost")
        my.cell(row=r, column=2, value=label)
        put(my, f"C{r}", 0, MONEY)
        year_cells(r, lambda k, orow=orow: f"{ovq}!$F${orow}*(1+CostEscalation)^{k - 1}")
    r += 1
    my.cell(row=r, column=1, value="Contingency")
    my.cell(row=r, column=2, value="Contingency on every cost")
    for col in "CDEFGH":
        put(my, f"{col}{r}", f'=SUMIF($A$5:$A${r - 1},"Cost",{col}5:{col}{r - 1})*Contingency', MONEY)
    put(my, f"I{r}", f"=SUM(C{r}:H{r})", MONEY)
    last = r
    add_table(my, "tblMultiYear", "A", 4, "I", last)
    labels = [("Total revenue", "Revenue"), ("Total costs", None), ("Net", None), ("Cumulative net", None)]
    rows_at = {}
    for label, _ in labels:
        r += 1
        rows_at[label] = r
        put(my, f"A{r}", label)
    for col in "CDEFGHI":
        rng = f"$A$5:$A${last}"
        put(my, f"{col}{rows_at['Total revenue']}", f'=SUMIF({rng},"Revenue",{col}5:{col}{last})', MONEY)
        put(my, f"{col}{rows_at['Total costs']}",
            f'=SUMIF({rng},"Cost",{col}5:{col}{last})+SUMIF({rng},"Contingency",{col}5:{col}{last})', MONEY)
        put(my, f"{col}{rows_at['Net']}", f"={col}{rows_at['Total revenue']}-{col}{rows_at['Total costs']}", MONEY)
    cum = rows_at["Cumulative net"]
    put(my, f"C{cum}", f"=C{rows_at['Net']}", MONEY)
    for prev, col in zip("CDEFG", "DEFGH"):
        put(my, f"{col}{cum}", f"={prev}{cum}+{col}{rows_at['Net']}", MONEY)
    put(my, f"I{cum}", f"=H{cum}", MONEY)
    total_row(my, rows_at["Total revenue"], 9)
    for label in ("Total costs", "Net", "Cumulative net"):
        my[f"A{rows_at[label]}"].font = BOLD
    r += 2
    put(my, f"A{r}", "Payback")
    periods = ["Build", *YEARS]
    formula = '"beyond the horizon"'
    for col, period in reversed(list(zip("CDEFGH", periods))):
        formula = f'IF({col}{cum}>=0,"{period}",{formula})'
    put(my, f"B{r}", "=" + formula)
    r += 1
    put(my, f"A{r}", "Return on cost (net / costs)")
    put(my, f"B{r}", f"=IF(I{rows_at['Total costs']}=0,\"\",I{rows_at['Net']}/I{rows_at['Total costs']})", "0.0%")
    widths(my, [14, 58, 16, 16, 16, 16, 16, 16, 18])
    my.freeze_panes = "C5"

    # --- Epics: deliverables, forecast and actual per epic
    ep = wb.create_sheet("Epics")
    put(ep, "A1", "Epics: deliverables, forecast and actual", font=TITLE)
    put(ep, "A2", f"In {ccy}. Forecast: each epic takes a share of the build in proportion to its planned working days (that "
                  "share of the delivery team and of the environments during delivery), plus the one-off costs tagged to it "
                  "and its Claude API budget; dates run the epics one after another from the delivery start date. Actual: "
                  "enter or keep the actual start and end dates; actual working days, the Claude API actual from the ledger, "
                  "and team and infra scaled by actual / planned days (same team and rates) unless you enter the timesheet "
                  "team cost. Blue cells are inputs; add an epic in a spare row.")
    header(ep, 4, ["Epic", "Name", "Deliverables", "Planned days", "Planned weeks", "Planned start", "Planned end",
                   "Share of build", "Team mandays", f"Team cost ({ccy})", f"Infra cost ({ccy})", f"Other costs ({ccy})",
                   f"Claude API budget ({ccy})", f"Forecast total ({ccy})", "Actual start", "Actual end", "Actual days",
                   "Days variance", "Days variance %", f"Claude API actual ({ccy})", f"Actual team cost, timesheets ({ccy})",
                   f"Actual total ({ccy})", f"Cost variance ({ccy})", "Cited from"])
    epics = model.get("epics") or []
    first_e, last_e = 5, 4 + len(epics) + SPARE_ROWS
    days_rng = f"$D${first_e}:$D${last_e}"
    env_delivery = "+".join(f"{ovq}!$D${row}" for label, row in run_rows if label.endswith("environment (run)")) or "0"
    oc = q("Other Costs")
    timeline = {t["name"]: t for t in (close or {}).get("timeline", []) if t.get("category") == "epic"}
    n_usage = len(close["usage"]) if close else 0
    detail = q("Claude Usage Detail")
    DATE = "yyyy-mm-dd"
    for i in range(len(epics) + SPARE_ROWS):
        r = first_e + i
        e_ = epics[i] if i < len(epics) else None
        if e_:
            put(ep, f"A{r}", e_["number"], "0", BLUE)
            ep.cell(row=r, column=2, value=e_["epic"])
            ep.cell(row=r, column=3, value="\n".join(e_.get("deliverables") or [])).alignment = WRAP
            put(ep, f"D{r}", e_["days"], QTY, BLUE)
            if e_.get("api_budget") is not None:
                put(ep, f"M{r}", e_["api_budget"], MONEY, BLUE)
            t = timeline.get(f"epic-{e_['number']}", {})
            a_start = e_.get("actual_start") or t.get("started")
            a_end = e_.get("actual_end") or (t.get("completed") if t.get("status", "done") == "done" else None)
            if a_start:
                put(ep, f"O{r}", iso_datetime(a_start), DATE, BLUE)
            if a_end:
                put(ep, f"P{r}", iso_datetime(a_end), DATE, BLUE)
            if e_.get("actual_team_cost") is not None:
                put(ep, f"U{r}", e_["actual_team_cost"], MONEY, BLUE)
            ep.cell(row=r, column=24, value=e_["cite"]).alignment = WRAP
        for col in "ABCDMOPU":
            if ep[f"{col}{r}"].value is None:
                ep[f"{col}{r}"].font = BLUE
                if col in "OP":
                    ep[f"{col}{r}"].number_format = DATE
        g = f'D{r}=""'  # a spare row stays blank until it has days
        put(ep, f"E{r}", f'=IF({g},"",D{r}/5)', "0.0")
        start = "ProjectStart" if i == 0 else f"WORKDAY(G{r - 1},1)"
        put(ep, f"F{r}", f'=IF(OR({g},ProjectStart=""),"",{start})', DATE)
        put(ep, f"G{r}", f'=IF(F{r}="","",WORKDAY(F{r},D{r}-1))', DATE)
        put(ep, f"H{r}", f'=IF({g},"",D{r}/SUM({days_rng}))', "0.0%")
        put(ep, f"I{r}", f'=IF({g},"",H{r}*{ovq}!$F${team_total})', QTY)
        put(ep, f"J{r}", f'=IF({g},"",H{r}*{ovq}!$G${team_total})', MONEY)
        put(ep, f"K{r}", f'=IF({g},"",H{r}*({env_delivery}))', MONEY)
        put(ep, f"L{r}", f'=IF(A{r}="","",SUMIFS({oc}!$C$5:$C${other_total - 1},{oc}!$F$5:$F${other_total - 1},A{r}))', MONEY)
        put(ep, f"N{r}", f'=IF({g},"",J{r}+K{r}+L{r}+N(M{r}))', MONEY)
        done = f'OR({g},O{r}="",P{r}="")'  # actuals only once both actual dates are in
        put(ep, f"Q{r}", f'=IF({done},"",NETWORKDAYS(O{r},P{r}))', QTY)
        put(ep, f"R{r}", f'=IF(Q{r}="","",Q{r}-D{r})', '+#,##0;-#,##0;0')
        put(ep, f"S{r}", f'=IF(Q{r}="","",Q{r}/D{r}-1)', '+0.0%;-0.0%;0.0%')
        if close:
            put(ep, f"T{r}", f'=IF(A{r}="","",SUMIFS({detail}!$P$5:$P${4 + max(n_usage, 1)},'
                             f'{detail}!$C$5:$C${4 + max(n_usage, 1)},A{r}&".*"))', MONEY)
        put(ep, f"V{r}", f'=IF(Q{r}="","",IF(U{r}="",J{r}*Q{r}/D{r},U{r})+K{r}*Q{r}/D{r}+L{r}+N(T{r}))', MONEY)
        put(ep, f"W{r}", f'=IF(V{r}="","",V{r}-N{r})', '+#,##0.00;-#,##0.00;0.00')
    add_table(ep, "tblEpics", "A", 4, "X", last_e)
    r = last_e + 1
    epic_total = r
    put(ep, f"A{r}", "Total")
    for col in "DIJKLMNQRTUVW":
        put(ep, f"{col}{r}", f"=SUM({col}{first_e}:{col}{last_e})", QTY if col in "DIQR" else MONEY)
    # variance over the epics that have actuals only
    put(ep, f"S{r}", f'=IF(Q{r}=0,"",Q{r}/SUMIFS(D{first_e}:D{last_e},Q{first_e}:Q{last_e},">0")-1)', '+0.0%;-0.0%;0.0%')
    total_row(ep, r, 24)
    r += 2
    put(ep, f"A{r}", f'=IF(D{epic_total}=DeliveryDays,"The planned epics add up to the days of delivery.",'
                     f'"The planned epics add up to "&D{epic_total}&" working days; the days of delivery input is "&DeliveryDays'
                     f'&". Costs are shared by planned epic days either way.")')
    widths(ep, [6, 26, 40, 9, 9, 12, 12, 9, 11, 14, 13, 13, 13, 15, 12, 12, 9, 9, 9, 13, 15, 15, 14, 40])
    ep.freeze_panes = "C5"

    # --- Sprints: fixed timeboxes of SprintDays working days from the delivery start, forecast and actual
    sp = wb.create_sheet("Sprints")
    put(sp, "A1", "Sprints: forecast and actual", font=TITLE)
    put(sp, "A2", f"In {ccy}. Every sprint is a timebox of the Overall input 'Working days per sprint' from the delivery start "
                  "date. Forecast: the sprint's planned working days (the days of delivery, sprint by sprint) at the delivery "
                  "team's and the environments' daily cost, plus any Claude API budget. Actual, up to the last actual activity "
                  "(the latest completed story or day of Claude usage): the days worked in the sprint, the stories completed, "
                  "the Claude API cost on its days, and the team at the days worked unless timesheets are entered. Sprints "
                  "past the plan appear when work ran over. Blue cells are inputs.")
    tl_rows = len((close or {}).get("timeline") or [])
    dy_rows = len((close or {}).get("daily") or [])
    tlq, dyq = q("Project Timeline"), q("Claude Usage Daily")
    last_parts = []
    if tl_rows:
        last_parts.append(f'MAXIFS({tlq}!$F$5:$F${4 + tl_rows},{tlq}!$A$5:$A${4 + tl_rows},"story")')
    if dy_rows:
        last_parts.append(f"MAX({dyq}!$A$5:$A${4 + dy_rows})")
    put(sp, "A3", "Last actual activity")
    if last_parts:
        put(sp, "C3", f'=IF(MAX({",".join(last_parts)})=0,"",MAX({",".join(last_parts)}))', "yyyy-mm-dd")
    wb.defined_names["LastActual"] = DefinedName("LastActual", attr_text="Sprints!$C$3")
    header(sp, 4, ["Sprint", "Start", "End", "Planned days", "Planned stories", f"Team cost ({ccy})", f"Infra cost ({ccy})",
                   f"Claude API budget ({ccy})", f"Forecast total ({ccy})", "Days worked", "Stories done",
                   "Stories done (list)", f"Claude API actual ({ccy})", f"Actual team cost, timesheets ({ccy})",
                   f"Actual total ({ccy})", f"Cost variance ({ccy})"])
    planned = {x["number"]: x for x in model.get("sprints") or []}
    sprint_figs = {x["number"]: x for x in (sprint_figures or [])}
    n_sprints = max(-(-int(a["delivery_days"]) // int(a["sprint_days"])), len(sprint_figs)) + SPARE_ROWS
    team_daily = f"{ovq}!$G${team_total}/DeliveryDays"
    infra_daily = f"({env_delivery})/DeliveryDays"
    first_s = 5
    for k in range(1, n_sprints + 1):
        r = first_s + k - 1
        put(sp, f"A{r}", k, "0")
        planned_days = f"MIN(SprintDays,MAX(0,DeliveryDays-({k}-1)*SprintDays))"
        timebox = f"WORKDAY(ProjectStart,({k}-1)*SprintDays)"
        # a sprint exists while it has planned days, or while actual work reaches into it
        exists = f'OR({planned_days}>0,AND(ProjectStart<>"",LastActual<>"",{timebox}<=N(LastActual)))'
        put(sp, f"D{r}", f'=IF({exists},{planned_days},"")', "0")
        put(sp, f"B{r}", f'=IF(OR(D{r}="",ProjectStart=""),"",{timebox})', "yyyy-mm-dd")
        put(sp, f"C{r}", f'=IF(B{r}="","",WORKDAY(B{r},SprintDays-1))', "yyyy-mm-dd")
        plan = planned.get(k, {})
        if plan.get("stories"):
            sp.cell(row=r, column=5, value="\n".join(plan["stories"])).alignment = WRAP
        sp[f"E{r}"].font = BLUE
        put(sp, f"F{r}", f'=IF(D{r}="","",D{r}*{team_daily})', MONEY)
        put(sp, f"G{r}", f'=IF(D{r}="","",D{r}*{infra_daily})', MONEY)
        if plan.get("api_budget") is not None:
            put(sp, f"H{r}", plan["api_budget"], MONEY, BLUE)
        else:
            sp[f"H{r}"].font = BLUE
        put(sp, f"I{r}", f'=IF(D{r}="","",F{r}+G{r}+N(H{r}))', MONEY)
        worked = f'OR(B{r}="",LastActual="",B{r}>N(LastActual))'  # no actuals before work reached the sprint
        put(sp, f"J{r}", f'=IF({worked},"",NETWORKDAYS(B{r},MIN(C{r},LastActual)))', "0")
        if tl_rows:
            put(sp, f"K{r}", f'=IF(J{r}="","",COUNTIFS({tlq}!$A$5:$A${4 + tl_rows},"story",'
                             f'{tlq}!$F$5:$F${4 + tl_rows},">="&B{r},{tlq}!$F$5:$F${4 + tl_rows},"<="&C{r}))', "0")
        done = sprint_figs.get(k, {}).get("stories_done_list")
        if done:
            sp.cell(row=r, column=12, value="\n".join(done)).alignment = WRAP
        if dy_rows:
            put(sp, f"M{r}", f'=IF(J{r}="","",SUMIFS({dyq}!$O$5:$O${4 + dy_rows},{dyq}!$A$5:$A${4 + dy_rows},">="&B{r},'
                             f'{dyq}!$A$5:$A${4 + dy_rows},"<="&C{r}))', MONEY)
        if plan.get("actual_team_cost") is not None:
            put(sp, f"N{r}", plan["actual_team_cost"], MONEY, BLUE)
        else:
            sp[f"N{r}"].font = BLUE
        put(sp, f"O{r}", f'=IF(J{r}="","",IF(N{r}="",J{r}*{team_daily},N{r})+J{r}*{infra_daily}+N(M{r}))', MONEY)
        put(sp, f"P{r}", f'=IF(O{r}="","",O{r}-N(I{r}))', '+#,##0.00;-#,##0.00;0.00')
    last_s = first_s + n_sprints - 1
    add_table(sp, "tblSprints", "A", 4, "P", last_s)
    r = last_s + 1
    put(sp, f"A{r}", "Total")
    for col in "DFGHIJKMNOP":
        put(sp, f"{col}{r}", f"=SUM({col}{first_s}:{col}{last_s})", "0" if col in "DJK" else MONEY)
    total_row(sp, r, 16)
    widths(sp, [8, 12, 12, 10, 34, 14, 13, 13, 15, 10, 10, 34, 14, 15, 15, 14])
    sp.freeze_panes = "B5"

    # --- Phase Budgets: the Claude API budget per planning phase against the ledger's actual
    pb = wb.create_sheet("Phase Budgets")
    put(pb, "A1", "Phase budgets: Claude API budget and actual per planning phase", font=TITLE)
    put(pb, "A2", f"In {ccy}, at Claude API list prices. The budget each planning phase may spend (test design, governance, "
                  "diagrams, course corrections ...), set once the epics and sprint plan are done; the epics and sprints carry "
                  "their own budgets on their sheets. Actual: the phase's Claude cost from the ledger (Claude Usage Detail) "
                  "once the close-out sheets are in. A control figure: not added to any total. Blue cells are inputs; "
                  "name a phase as the token ledger does.")
    header(pb, 4, ["Phase", f"Claude API budget ({ccy})", f"Claude API actual ({ccy})", f"Variance ({ccy})", "Variance %",
                   "Cited from"])
    budgets = model.get("phase_budgets") or []
    first_b, last_b = 5, 4 + len(budgets) + SPARE_ROWS
    for i in range(len(budgets) + SPARE_ROWS):
        r = first_b + i
        b_ = budgets[i] if i < len(budgets) else None
        if b_:
            put(pb, f"A{r}", b_["phase"], None, BLUE)
            put(pb, f"B{r}", b_["api_budget"], MONEY, BLUE)
            pb.cell(row=r, column=6, value=b_["cite"]).alignment = WRAP
        for col in "AB":
            pb[f"{col}{r}"].font = BLUE
        if close:
            put(pb, f"C{r}", f'=IF(A{r}="","",SUMIFS({detail}!$P$5:$P${4 + max(n_usage, 1)},'
                             f'{detail}!$B$5:$B${4 + max(n_usage, 1)},A{r}))', MONEY)
        put(pb, f"D{r}", f'=IF(OR(A{r}="",C{r}=""),"",C{r}-N(B{r}))', '+#,##0.00;-#,##0.00;0.00')
        put(pb, f"E{r}", f'=IF(OR(D{r}="",N(B{r})=0),"",C{r}/B{r}-1)', '+0.0%;-0.0%;0.0%')
    add_table(pb, "tblPhaseBudgets", "A", 4, "F", last_b)
    r = last_b + 1
    put(pb, f"A{r}", "Total")
    for col in "BCD":
        put(pb, f"{col}{r}", f"=SUM({col}{first_b}:{col}{last_b})", MONEY)
    total_row(pb, r, 6)
    widths(pb, [34, 16, 16, 14, 11, 40])
    pb.freeze_panes = "B5"

    # --- Sources & Gaps
    sg = wb.create_sheet("Sources & Gaps")
    put(sg, "A1", "Sources, assumptions and gaps", font=TITLE)
    lines: list[tuple[str, str]] = [("Price basis", ""),
                                    ("", f"Azure Retail Prices API list prices (pay-as-you-go, no discounts), region {region}, "
                                         f"{cloud_ccy}, fetched {price_date}.")]
    if cloud_ccy != ccy:
        lines.append(("", f"Exchange rate 1 {cloud_ccy} = {a['fx_rate']} {ccy}: {model.get('fx_source', '')}"))
    lines.append(("", f"Rate card: {card.get('source', 'not stated')}, updated {card.get('updated', 'not stated')}."))
    if close:
        pr = close["prices"]
        lines.append(("", f"Claude API list prices from {pr.get('source', '')}, read {pr.get('fetched', '')}, USD per million "
                          "tokens. Claude Code on a subscription is billed per seat: this is the API-equivalent cost."))
        if close["unpriced"]:
            lines.append(("", "Claude models with no price on the page (their tokens are not costed): " + ", ".join(close["unpriced"])))
    lines.append(("", f"Horizon: {a['horizon_years']} years of run after go-live, run costs rising {a['cost_escalation_pct']}% a year."))
    lines += [("", ""), ("Gaps: not priced and not in any total", "")]
    gaps = list(model.get("gaps") or [])
    if not model.get("revenue"):
        gaps.append("Revenue: none given, so the Multi-Year sheet shows costs only.")
    lines += [("", g) for g in gaps] or [("", "None recorded.")]
    assumed = []
    for env in model["environments"]:
        for item in env["items"]:
            kind, _, ref = split_cite(item["cite"])
            if kind != "file":
                assumed.append(f"{env['name']} / {item['component']}: {kind}: {ref}")
    for m in model.get("team") or []:
        kind, _, ref = split_cite(m["cite"])
        if kind != "file":
            assumed.append(f"Team / {m['role']}: {kind}: {ref}")
    for label, entries, key in (("Other", model.get("other_costs"), "item"), ("Run cost", model.get("run_costs"), "item"),
                                ("Support", model.get("support"), "role"), ("Revenue", model.get("revenue"), "stream"),
                                ("Phase budget", model.get("phase_budgets"), "phase")):
        for o in entries or []:
            kind, _, ref = split_cite(o["cite"])
            if kind != "file":
                assumed.append(f"{label} / {o[key]}: {kind}: {ref}")
    lines += [("", ""), ("Assumptions and owner statements (not from the BMad artifacts)", "")]
    lines += [("", t) for t in assumed] or [("", "None.")]
    lines += [("", ""), (f"Stamp: built {stamp['built']}; sha256 of each cited file", "")]
    lines += [("", f"{path}  {digest[:12]}") for path, digest in stamp["files"].items()]
    for i, (head, text) in enumerate(lines, 3):
        if head:
            put(sg, f"A{i}", head, font=SECTION)
        if text:
            sg.cell(row=i, column=2, value=text).alignment = Alignment(wrap_text=False)
    widths(sg, [4, 120])

    if close:
        write_close_out(wb, close, ccy, dict(put=put, header=header, widths=widths, total_row=total_row,
                                             add_table=add_table, q=q, TITLE=TITLE, SECTION=SECTION, MONEY=MONEY,
                                             QTY=QTY, BLUE=BLUE, GREEN=GREEN, BOLD=BOLD))
    order = ["Overall", "Multi-Year", "Epics", "Sprints", "Phase Budgets", "Revenue", "Rate Card", "Other Costs", "Run Costs",
             *[e["name"] for e in model["environments"]],
             "Claude Prices", "Claude Usage", "Claude Usage Detail", "Claude Usage Daily", "Project Timeline", "AI Feedback",
             "Sources & Gaps"]
    for i, name in enumerate(n for n in order if n in wb.sheetnames):
        wb.move_sheet(name, offset=i - wb.sheetnames.index(name))
    wb.calculation.fullCalcOnLoad = True
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return wb.sheetnames


def claude_models(close: dict) -> list[str]:
    return sorted({r["model"] for r in close["usage"]})


def iso_datetime(value: str):
    """'2026-09-27T11:02:59.394Z' or '2026-09-27' -> a naive datetime Excel and Power BI read as a date."""
    if not value:
        return None
    try:
        return datetime.fromisoformat(value.replace("Z", "+00:00")).replace(tzinfo=None)
    except ValueError:
        return value


def write_close_out(wb, close: dict, ccy: str, k: dict) -> None:
    put, header, widths, add_table, q = k["put"], k["header"], k["widths"], k["add_table"], k["q"]
    MONEY, QTY, BLUE = k["MONEY"], k["QTY"], k["BLUE"]
    DATE, TOKENS = "yyyy-mm-dd hh:mm", "#,##0"
    kinds = ("input", "cache_write_5m", "cache_write_1h", "cache_read", "output")

    # Claude Prices
    pr = close["prices"]
    ps = wb.create_sheet("Claude Prices")
    put(ps, "A1", "Claude API list prices", font=k["TITLE"])
    put(ps, "A2", f"USD per million tokens, from {pr.get('source', '')}, read {pr.get('fetched', '')}. Blue cells are inputs.")
    header(ps, 4, ["Pricing model", "Name", "Input", "Cache write 5m", "Cache write 1h", "Cache read", "Output"])
    r = 4
    for mid, price in sorted(pr.get("models", {}).items()):
        r += 1
        ps.cell(row=r, column=1, value=mid)
        ps.cell(row=r, column=2, value=price.get("name", mid))
        for c, key in enumerate(kinds, 3):
            put(ps, f"{chr(64 + c)}{r}", price[key], "#,##0.00", BLUE)
    add_table(ps, "tblClaudePrices", "A", 4, "G", r)
    price_range = f"{q('Claude Prices')}!$A$5:$G${max(r, 5)}"
    widths(ps, [24, 22, 12, 14, 14, 12, 12])

    # Claude Usage Detail: one row per work item x model
    ds = wb.create_sheet("Claude Usage Detail")
    put(ds, "A1", "Claude API calls by work item and model", font=k["TITLE"])
    put(ds, "A2", "From the token ledger (claude_usage.csv). Cost = tokens x the Claude Prices sheet; models with no price cost 0.")
    header(ds, 4, ["Kind", "Work item", "Story no", "Title", "Model", "Pricing model", "First call", "Last call", "Calls",
                   "Input tokens", "Cache write 5m tokens", "Cache write 1h tokens", "Cache read tokens", "Output tokens",
                   "Cost (USD)", f"Cost ({ccy})"])
    r = 4
    for u in close["usage"]:
        r += 1
        values = [u["kind"], u["work_item"], u.get("story_no", ""), u.get("title", ""), u["model"], pricing_id(u["model"]),
                  iso_datetime(u.get("first_call", "")), iso_datetime(u.get("last_call", "")), int(u["calls"])]
        values += [int(u.get(f"{key}_tokens") or 0) for key in kinds]
        for c, v in enumerate(values, 1):
            cell = ds.cell(row=r, column=c, value=v)
            cell.number_format = DATE if c in (7, 8) else TOKENS if c >= 9 else "General"
        vl = [f"J{r}*VLOOKUP($F{r},{price_range},3,FALSE)", f"K{r}*VLOOKUP($F{r},{price_range},4,FALSE)",
              f"L{r}*VLOOKUP($F{r},{price_range},5,FALSE)", f"M{r}*VLOOKUP($F{r},{price_range},6,FALSE)",
              f"N{r}*VLOOKUP($F{r},{price_range},7,FALSE)"]
        put(ds, f"O{r}", f"=IFERROR(({'+'.join(vl)})/1000000,0)", MONEY)
        put(ds, f"P{r}", f"=O{r}*ClaudeFxRate", MONEY)
    add_table(ds, "tblClaudeUsageDetail", "A", 4, "P", r)
    last_detail = max(r, 5)
    widths(ds, [10, 30, 9, 30, 26, 22, 17, 17, 9, 14, 14, 14, 16, 14, 14, 14])
    ds.freeze_panes = "A5"

    # Claude Usage: by model, totals
    us = wb.create_sheet("Claude Usage")
    put(us, "A1", "Claude API usage", font=k["TITLE"])
    put(us, "A2", "API calls and tokens counted from this project's Claude Code logs, costed at Claude API list prices. "
                  "On a Pro, Max, Team or Enterprise subscription the bill is per seat: this is the API-equivalent cost.")
    header(us, 4, ["Model", "Calls", "Input tokens", "Cache write 5m tokens", "Cache write 1h tokens", "Cache read tokens",
                   "Output tokens", "Cost (USD)", f"Cost ({ccy})"])
    detail = q("Claude Usage Detail")
    r = 4
    for m in claude_models(close):
        r += 1
        us.cell(row=r, column=1, value=m)
        for c, col in zip(range(2, 10), "IJKLMNOP"):
            # "="&A forces an exact match: a bare model name such as <synthetic> would read as a comparison
            put(us, f"{chr(64 + c)}{r}", f'=SUMIFS({detail}!${col}$5:${col}${last_detail},{detail}!$E$5:$E${last_detail},"="&$A{r})',
                MONEY if c >= 8 else TOKENS)
    add_table(us, "tblClaudeUsage", "A", 4, "I", r)
    r += 1
    put(us, f"A{r}", "Total")
    for c in range(2, 10):
        col = chr(64 + c)
        put(us, f"{col}{r}", f"=SUM({col}5:{col}{r - 1})", MONEY if c >= 8 else TOKENS)
    k["total_row"](us, r, 9)
    widths(us, [28, 10, 14, 16, 16, 16, 14, 14, 14])

    # Claude Usage Daily: one row per day x work item x model, for day-to-day dashboards
    if close["daily"]:
        dy = wb.create_sheet("Claude Usage Daily")
        put(dy, "A1", "Claude API usage by day", font=k["TITLE"])
        put(dy, "A2", "From the token ledger (claude_usage_daily.csv, refreshed after every turn). Cost = tokens x the Claude "
                      "Prices sheet; models with no price cost 0.")
        header(dy, 4, ["Day", "Kind", "Work item", "Story no", "Title", "Model", "Pricing model", "Calls", "Input tokens",
                       "Cache write 5m tokens", "Cache write 1h tokens", "Cache read tokens", "Output tokens",
                       "Cost (USD)", f"Cost ({ccy})"])
        r = 4
        for u in close["daily"]:
            r += 1
            values = [iso_datetime(u["day"]), u["kind"], u["work_item"], u.get("story_no", ""), u.get("title", ""),
                      u["model"], pricing_id(u["model"]), int(u["calls"])]
            values += [int(u.get(f"{key}_tokens") or 0) for key in kinds]
            for c, v in enumerate(values, 1):
                dy.cell(row=r, column=c, value=v).number_format = "yyyy-mm-dd" if c == 1 else TOKENS if c >= 8 else "General"
            vl = "+".join(f"{col}{r}*VLOOKUP($G{r},{price_range},{n},FALSE)" for col, n in zip("IJKLM", range(3, 8)))
            put(dy, f"N{r}", f"=IFERROR(({vl})/1000000,0)", MONEY)
            put(dy, f"O{r}", f"=N{r}*ClaudeFxRate", MONEY)
        add_table(dy, "tblClaudeUsageDaily", "A", 4, "O", r)
        widths(dy, [12, 10, 30, 9, 30, 26, 22, 9, 14, 14, 14, 16, 14, 14, 14])
        dy.freeze_panes = "A5"

    # Project Timeline
    if close["timeline"]:
        ts = wb.create_sheet("Project Timeline")
        put(ts, "A1", "Project timeline", font=k["TITLE"])
        project = next((t for t in close["timeline"] if t["category"] == "project"), {})
        put(ts, "A2", f"Started {project.get('started', '')}; "
                      + (f"completed {project['completed']}" if project.get("completed") else "still running")
                      + f"; {project.get('days', '')} calendar days, {project.get('detail', '')}. "
                      "From git history of the BMad artifacts and sprint status, the agents' sanctums and the token ledger.")
        header(ts, 4, ["Category", "Name", "Detail", "Status", "Started", "Completed", "Days", "Source"])
        r = 4
        for t in close["timeline"]:
            r += 1
            values = [t["category"], t["name"], t.get("detail", ""), t.get("status", ""), iso_datetime(t.get("started", "")),
                      iso_datetime(t.get("completed", "")), int(t["days"]) if str(t.get("days", "")).isdigit() else None,
                      t.get("source", "")]
            for c, v in enumerate(values, 1):
                cell = ts.cell(row=r, column=c, value=v)
                if c in (5, 6):
                    cell.number_format = "yyyy-mm-dd"
        add_table(ts, "tblTimeline", "A", 4, "H", r)
        widths(ts, [12, 36, 30, 12, 12, 12, 8, 14])
        ts.freeze_panes = "A5"

    # AI Feedback: the team's rating of the AI at the end of each epic
    if close["feedback"]:
        fb = wb.create_sheet("AI Feedback")
        put(fb, "A1", "AI feedback by epic", font=k["TITLE"])
        put(fb, "A2", "The team's rating of the AI at the end of each epic, 1 (very poor) to 5 (excellent), and what went "
                      "wrong when it was poor (1 or 2). From ai-feedback.csv in Scrooge's sanctum.")
        header(fb, 4, ["Epic", "Rating", "Poor", "Categories", "Other", "Comment", "Status", "Recorded"])
        r = 4
        for f in close["feedback"]:
            r += 1
            values = [f["epic"], int(f["rating"]) if str(f.get("rating", "")).isdigit() else None, f.get("poor", ""),
                      f.get("categories", ""), f.get("other", ""), f.get("comment", ""), f.get("status", ""),
                      iso_datetime(f.get("recorded", ""))]
            for c, v in enumerate(values, 1):
                cell = fb.cell(row=r, column=c, value=v)
                if c == 8:
                    cell.number_format = "yyyy-mm-dd"
        add_table(fb, "tblAIFeedback", "A", 4, "H", r)
        r += 1
        put(fb, f"A{r}", "Average")
        put(fb, f"B{r}", f"=IFERROR(AVERAGE(B5:B{r - 1}),\"\")", "0.0")
        k["total_row"](fb, r, 8)
        widths(fb, [10, 8, 7, 60, 36, 36, 10, 12])
        fb.freeze_panes = "A5"


def write_facts(model: dict, figures: dict, out: Path) -> None:
    """Every revenue and cost line per period as a plain value, one row each, for dashboards (Power BI, Excel Online).

    Long format: estimate, currency, period (Build, Year 1..), year (0 for the build), category, item, amount.
    """
    a = settings(model)
    ccy, name, h, e, cont = model["currency"], model["name"], a["horizon_years"], a["escalation"], a["contingency"]
    rows = [["estimate", "currency", "period", "year", "category", "item", "amount"]]

    def add(year: int, category: str, item: str, amount: float) -> None:
        rows.append([name, ccy, "Build" if year == 0 else f"Year {year}", year, category, item, round(amount, 2)])

    build = [("delivery team", "Delivery team", figures["team"]["cost"]), ("one-off", "Other one-off costs", figures["other_costs"])]
    build += [("environment", f"{x['name']} environment (during delivery)", x["delivery"]) for x in figures["environments"]]
    claude = figures.get("claude")
    if claude and claude["in_totals"]:
        build.append(("claude", "Claude API usage", claude["cost"]))
    for category, item, amount in build:
        add(0, category, item, amount)
    add(0, "contingency", "Contingency", sum(x[2] for x in build) * cont)
    if claude and not claude["in_totals"]:
        add(0, "claude (memo)", "Claude API usage at API list prices (not in the totals)", claude["cost"])
    for k in range(1, h + 1):
        factor = (1 + e) ** (k - 1)
        for stream in model.get("revenue") or []:
            add(k, "revenue", stream["stream"], stream["amounts"][k - 1] if k <= len(stream["amounts"]) else 0)
        run = [("environment", f"{x['name']} environment (run)", x["run_per_year"] * factor) for x in figures["environments"]]
        run += [("support", "Run support (OPEX)", figures["support"]["cost_per_year"] * factor),
                ("run costs", "Recurring run costs", figures["run_costs_per_year"] * factor)]
        for category, item, amount in run:
            add(k, category, item, amount)
        add(k, "contingency", "Contingency", sum(x[2] for x in run) * cont)
    with out.open("w", newline="", encoding="utf-8") as fh:
        csv.writer(fh).writerows(rows)


def write_epics(model: dict, figures: dict, out: Path) -> None:
    """One row per epic: deliverables, duration and cost, as plain values for dashboards."""
    fields = ["number", "epic", "days", "weeks", "start", "end", "share", "mandays", "team_cost", "infra_cost",
              "other_costs", "api_budget", "forecast_total", "actual_start", "actual_end", "actual_days", "days_variance",
              "days_variance_pct", "api_actual", "actual_team_cost", "actual_total", "cost_variance"]
    deliverables = {x["number"]: "; ".join(x.get("deliverables") or []) for x in model.get("epics") or []}
    with out.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["estimate", "currency", *fields[:2], "deliverables", *fields[2:]])
        for x in figures["epics"]:
            w.writerow([model["name"], model["currency"], x["number"], x["epic"], deliverables.get(x["number"], ""),
                        *[("" if x[f] is None else x[f]) for f in fields[2:]]])


def write_sprints(model: dict, figures: dict, out: Path) -> None:
    """One row per sprint: dates, forecast and actual, as plain values for dashboards."""
    fields = ["number", "start", "end", "days", "team_cost", "infra_cost", "api_budget", "forecast_total", "days_worked",
              "stories_done", "api_actual", "actual_team_cost", "actual_total", "cost_variance"]
    with out.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["estimate", "currency", *fields, "planned_stories", "stories_done_list"])
        for x in figures["sprints"]:
            w.writerow([model["name"], model["currency"], *[("" if x[f] is None else x[f]) for f in fields],
                        "; ".join(x["planned_stories"]), "; ".join(x["stories_done_list"])])


def infra_rows(model: dict) -> list[dict]:
    """Infrastructure OPEX after go-live: each priced cloud line, each environment's total, then the grand total."""
    a = settings(model)
    rows, grand = [], {"daily": 0.0, "weekly": 0.0, "monthly": 0.0, "annual": 0.0}
    for env in model["environments"]:
        months = env.get("run_months_per_year", 12)
        total = dict.fromkeys(grand, 0.0)
        for item in env["items"]:
            monthly = item["unit_price"] * item["quantity"] * item_usage(item, a) * a["fx_rate"]
            daily = monthly * 12 / 365
            cost = {"daily": daily, "weekly": daily * 7, "monthly": monthly, "annual": monthly * months}
            rows.append({"line": "item", "environment": env["name"], "component": item["component"],
                         "service": item["service"], "sku": item.get("sku", ""), "run_months_per_year": months, **cost,
                         "cite": item["cite"], "price_source": item["price_source"]})
            for k in total:
                total[k] += cost[k]
        rows.append({"line": "environment total", "environment": env["name"], "component": "", "service": "", "sku": "",
                     "run_months_per_year": months, **total, "cite": "", "price_source": ""})
        for k in grand:
            grand[k] += total[k]
    rows.append({"line": "total", "environment": "", "component": "", "service": "", "sku": "", "run_months_per_year": "",
                 **grand, "cite": "", "price_source": ""})
    return rows


def write_infra_cost(model: dict, out: Path) -> None:
    """The infrastructure OPEX (daily, weekly, monthly, annual) as plain values, for people and dashboards."""
    fields = ["line", "environment", "component", "service", "sku", "run_months_per_year", "daily", "weekly", "monthly",
              "annual", "cite", "price_source"]
    with out.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["estimate", "currency", *fields])
        for x in infra_rows(model):
            w.writerow([model["name"], model["currency"],
                        *[round(x[f], 2) if isinstance(x[f], float) else x[f] for f in fields]])


def write_phases(model: dict, figures: dict, out: Path) -> None:
    """One row per budgeted planning phase: Claude API budget against actual, as plain values for dashboards."""
    fields = ["phase", "api_budget", "api_actual", "variance", "variance_pct"]
    with out.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["estimate", "currency", *fields])
        for x in figures["phase_budgets"]:
            w.writerow([model["name"], model["currency"], *[("" if x[f] is None else x[f]) for f in fields]])


# --------------------------------------------------------------------------- main


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    p.add_argument("project_root")
    p.add_argument("model")
    p.add_argument("--rate-card", type=Path)
    p.add_argument("--out", type=Path)
    p.add_argument("--check", action="store_true", help="validate only")
    p.add_argument("--drift", action="store_true", help="report cited files changed since the build")
    p.add_argument("--force", action="store_true", help="overwrite a workbook edited after the last build")
    args = p.parse_args(sys.argv[1:] if argv is None else argv)

    root = Path(args.project_root).resolve()
    model_path = Path(args.model)
    if not model_path.is_absolute():
        model_path = (root / model_path) if (root / model_path).is_file() else estimates_folder(root) / model_path
    try:
        model = json.loads(model_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as e:
        print(json.dumps({"status": "error", "message": f"cannot read model {model_path}: {e}"}))
        return 2

    if args.drift:
        result = drift(root, model)
        print(json.dumps({"model": str(model_path), **result}, indent=1))
        return 1 if result["status"] == "drift" else 0

    card, card_path = load_rate_card(root, model, args.rate_card)
    errors = validate(root, model, card)
    close, close_errors = load_close_out(root, model)
    errors += close_errors
    if model.get("team") and not card:
        errors.append(f"no rate card: write {card_path} (roles and day rates) or add 'rate_card' to the model")
    if errors:
        print(json.dumps({"status": "invalid", "model": str(model_path), "errors": errors}, indent=1))
        return 1
    figures = compute(model, card, close)
    if args.check:
        print(json.dumps({"status": "valid", "model": str(model_path), "figures": figures}, indent=1))
        return 0

    out = args.out or model_path.with_name(model_path.name.removesuffix(".estimate.json").removesuffix(".json") + ".xlsx")
    if not out.resolve().is_relative_to(root):
        print(json.dumps({"status": "error", "message": f"workbook must be inside the project: {out}"}))
        return 2
    saved = model.get("stamp", {}).get("workbook_mtime")
    if out.is_file() and not args.force:
        if saved is None or abs(out.stat().st_mtime - saved) > 1:
            print(json.dumps({"status": "error", "workbook": str(out),
                              "message": "the workbook was changed after the last build (by hand?); rebuilding "
                                         "loses those changes. Carry them into the model, or pass --force"}))
            return 2
    stamp = make_stamp(root, model, card_path, close["files"] if close else [])
    try:
        sheets = write_workbook(model, card, stamp, out, close, figures.get("sprints"))
        facts = out.with_name(out.stem + ".facts.csv")
        write_facts(model, figures, facts)
        if figures["epics"]:
            write_epics(model, figures, out.with_name(out.stem + ".epics.csv"))
        write_sprints(model, figures, out.with_name(out.stem + ".sprints.csv"))
        infra = out.with_name(out.stem + ".infra_cost.csv")
        write_infra_cost(model, infra)
        if figures["phase_budgets"]:
            write_phases(model, figures, out.with_name(out.stem + ".phases.csv"))
    except PermissionError:
        print(json.dumps({"status": "error", "message": f"cannot write {out}: is it open in Excel?"}))
        return 2
    model["stamp"] = {**stamp, "workbook_mtime": out.stat().st_mtime}
    model_path.write_text(json.dumps(model, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({"status": "ok", "workbook": str(out), "facts": str(facts), "infra_cost": str(infra),
                      "model": str(model_path), "sheets": sheets,
                      "figures": figures}, indent=1))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
