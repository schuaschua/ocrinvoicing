#!/usr/bin/env python3
# /// script
# requires-python = ">=3.11"
# dependencies = ["python-docx>=1.1"]
# ///
"""Unit tests for build-docx.py (stdlib unittest; python-docx to read the documents back).

Run from the skill root: uv run scripts/tests/test_build-docx.py
"""

from __future__ import annotations

import json
import os
import struct
import sys
import unittest
import zlib
from pathlib import Path

import docx

sys.path.insert(0, str(Path(__file__).resolve().parent))
from arch_fixtures import *  # noqa: E402,F403

OUTLINE = [line[3:].strip() for line in (SCRIPTS.parent / "assets/design-docs/hld.md").read_text().splitlines()
           if line.startswith("## ")]


def tiny_png(path: Path) -> None:
    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))
    path.write_bytes(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
                     + chunk(b"IDAT", zlib.compress(b"\x00\xff\xff\xff")) + chunk(b"IEND", b""))


def hld_model(**extra) -> dict:
    sections = [{"heading": h, "level": 1, "blocks": [{"gap": f"{h}: not stated."}]} for h in OUTLINE]
    sections[0]["blocks"] = [{"text": "A chat service that answers customers' questions.", "cite": ["SPEC.md §2", "AD-1"]}]
    sections[3]["blocks"] = [
        {"table": {"columns": ["Container", "Technology"], "rows": [["Chat API", ".NET 8"], ["Store", "Cosmos DB"]]},
         "cite": ["AD-1", "AD-2"]},
        {"diagram": "c4-containers"}]
    sections.insert(4, {"heading": "Answers", "level": 2, "blocks": [{"bullets": ["Azure OpenAI answers"], "cite": "AD-3"}]})
    model = {"type": "hld", "title": "Chat service", "version": "0.2", "status": "Draft", "author": "Tester",
             "spine": SPINE_REL, "sources": ["_bmad-output/specs/SPEC.md"], "sections": sections,
             "gaps": ["No AD states the Azure region."]}
    model.update(extra)
    return model


class DocxTests(Project):
    def setUp(self):
        super().setUp()
        self.docs = self.root / "docs/architecture/design"
        self.docs.mkdir(parents=True)
        self.assertEqual(self.build(containers_model())[0], 0)

    def doc(self, model: dict, *flags: str, name: str = "hld") -> tuple[int, dict]:
        path = self.docs / f"{name}.doc.json"
        if path.is_file():  # an edited model keeps the stamp of its last build
            model = {**model, "stamp": json.loads(path.read_text(encoding="utf-8")).get("stamp")}
        path.write_text(json.dumps(model), encoding="utf-8")
        return run("build-docx.py", str(self.root), str(path), *flags)

    def paragraphs(self, name: str = "hld") -> list[tuple[str, str]]:
        d = docx.Document(str(self.docs / f"{name}.docx"))
        return [(p.style.name, p.text) for p in d.paragraphs if p.text.strip()]

    def test_builds_numbered_cited_document(self):
        code, out = self.doc(hld_model())
        self.assertEqual(code, 0, out)
        paras = self.paragraphs()
        headings = [t for s, t in paras if s.startswith("Heading")]
        self.assertEqual(headings[0], "1 Introduction")
        self.assertIn("4.1 Answers", headings)
        self.assertIn("Appendix A: Open Points", headings)
        self.assertTrue(any("[SPEC.md §2; AD-1]" in t for _, t in paras))
        self.assertTrue(any(t.startswith("Figure 1: Chat service") for _, t in paras))
        self.assertEqual(out["uncited_ads"], [])
        self.assertEqual(out["open_points"], len(OUTLINE) - 2 + 1)
        self.assertEqual(out["diagrams_without_png"], ["c4-containers"])
        stamp = json.loads((self.docs / "hld.doc.json").read_text())["stamp"]
        self.assertEqual((stamp["arch_type"], stamp["cited_ads"]), ("hld", "AD-1,AD-2,AD-3"))

    def test_embeds_the_png_export(self):
        tiny_png(self.folder / "c4-containers.png")
        code, out = self.doc(hld_model())
        self.assertEqual((code, out["diagrams_without_png"]), (0, []), out)
        self.assertEqual(len(docx.Document(str(self.docs / "hld.docx")).inline_shapes), 1)

    def test_refuses_what_it_cannot_trace(self):
        cases = [
            (lambda m: m["sections"].pop(2), "outline section 'Architecture Overview' is missing"),
            (lambda m: m["sections"][0]["blocks"][0].update(cite=["AD-9"]), "does not define"),
            (lambda m: m["sections"][0]["blocks"][0].update(cite=["SPEC.md §7"]), "names no section of SPEC.md"),
            (lambda m: m["sections"][0]["blocks"][0].pop("cite"), "make it a gap block"),
            (lambda m: m["sections"][3]["blocks"][1].update(diagram="deployment"), "draw it first"),
            (lambda m: m.update(sources=["docs/notes.md"]), "outside _bmad-output"),
            (lambda m: m.update(type="lld"), "scope is required"),
        ]
        for change, message in cases:
            model = hld_model()
            change(model)
            code, out = self.doc(model, "--check")
            self.assertEqual(code, 1, message)
            self.assertIn(message, out["message"])

    def test_refuses_a_stale_diagram_and_reports_document_drift(self):
        self.assertEqual(self.doc(hld_model())[0], 0)
        spine = self.root / SPINE_REL
        spine.write_text(SPINE.replace(".NET 8", ".NET 9"), encoding="utf-8")
        code, out = run("diagram-drift.py", str(self.root))
        self.assertEqual(code, 1)
        self.assertEqual(out["documents"][0]["status"], "stale")
        self.assertEqual(out["documents"][0]["ads"][0]["id"], "AD-1")
        code, out = self.doc(hld_model(), "--check")
        self.assertIn("is stale; redraw it", out["message"])

    def test_keeps_word_edits_unless_forced(self):
        self.assertEqual(self.doc(hld_model())[0], 0)
        self.assertEqual(self.doc(hld_model())[0], 0)  # untouched since the build: replaced
        target = self.docs / "hld.docx"
        later = target.stat().st_mtime + 60
        os.utime(target, (later, later))
        code, out = self.doc(hld_model())
        self.assertEqual(code, 1)
        self.assertIn("changed since the last build", out["message"])
        self.assertEqual(self.doc(hld_model(), "--force")[0], 0)

    def test_confluence_markdown_matches_the_word_build(self):
        tiny_png(self.folder / "c4-containers.png")
        code, out = self.doc(hld_model(), "--formats", "docx,confluence")
        self.assertEqual(code, 0, out)
        page = out["confluence"]
        md = Path(page["file"]).read_text(encoding="utf-8")
        self.assertEqual(page["title"], "Chat service High-Level Design")
        self.assertEqual(page["attachments"], [str(self.folder / "c4-containers.png")])
        self.assertIn("# 1 Introduction", md)
        self.assertIn("## 4.1 Answers", md)
        self.assertIn("_[SPEC.md §2; AD-1]_", md)
        self.assertIn("> **Open point 1:**", md)
        self.assertIn("| Store | Cosmos DB |\n\n_[AD-1; AD-2]_", md)
        self.assertIn("# Appendix B: Traceability", md)
        word = [t for _, t in self.paragraphs() if t[0].isdigit() and " " in t]
        for heading in word[:3]:
            self.assertIn(heading, md)
        # confluence alone leaves the Word file and its edit guard alone
        target = self.docs / "hld.docx"
        later = target.stat().st_mtime + 60
        os.utime(target, (later, later))
        self.assertEqual(self.doc(hld_model(), "--formats", "confluence")[0], 0)
        self.assertEqual(self.doc(hld_model())[0], 1)

    def test_project_outline_and_word_template(self):
        templates = self.docs / "templates"
        templates.mkdir()
        (templates / "lld.md").write_text("# LLD\n\n## Scope\nWhat is covered.\n\n## API\nEndpoints.\n", encoding="utf-8")
        ref = docx.Document()
        ref.sections[0].header.paragraphs[0].text = "ACME CONFIDENTIAL"
        ref.add_paragraph("template body text")
        ref.save(str(templates / "reference.docx"))
        model = {"type": "lld", "title": "Chat service", "scope": "Chat API", "spine": SPINE_REL, "sources": [],
                 "sections": [{"heading": "Scope", "blocks": [{"text": "The Chat API.", "cite": "AD-1"}]},
                              {"heading": "API", "blocks": [{"gap": "No endpoints are stated."}]}]}
        code, out = self.doc(model, name="lld-chat-api")
        self.assertEqual(code, 0, out)
        self.assertTrue(out["outline"].endswith("templates/lld.md"))
        d = docx.Document(str(self.docs / "lld-chat-api.docx"))
        self.assertEqual(d.sections[0].header.paragraphs[0].text, "ACME CONFIDENTIAL")
        texts = [p.text for p in d.paragraphs]
        self.assertNotIn("template body text", texts)
        self.assertIn("Low-Level Design: Chat API", texts)


if __name__ == "__main__":
    unittest.main()
