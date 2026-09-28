#!/usr/bin/env python3
# /// script
# requires-python = ">=3.11"
# dependencies = ["python-docx>=1.1"]
# ///
"""Build a cited High-Level or Low-Level Design Word document from its model.

The model, <name>.doc.json in the design documents folder (org config design_docs_folder,
default docs/architecture/design), holds the document's sections in the order of its outline,
each made of blocks that cite the spine's ADs or a section of a BMad source:

  {"type": "hld" | "lld", "title": "Burdeebur", "scope": "Back-end API" (LLD), "version": "0.1",
   "status": "Draft", "author": "Alex Tan",
   "spine": "_bmad-output/planning-artifacts/architecture/<name>/ARCHITECTURE-SPINE.md",
   "sources": ["_bmad-output/specs/SPEC.md", ...],
   "sections": [{"heading": "Introduction", "level": 1, "blocks": [
       {"text": "...", "cite": ["SPEC.md §1"]},
       {"bullets": ["...", "..."], "cite": ["AD-2"]},
       {"table": {"columns": ["Component", "Technology"], "rows": [["API", ".NET 9"]]}, "cite": ["AD-3"]},
       {"diagram": "c4-context"},
       {"gap": "The spec does not say who the audience is."}]}],
   "gaps": ["document-wide open points"],
   "confluence": {...}}                  the page it was published to or started from (see references/confluence.md;
                                         kept by Da Vinci, ignored by the build)

A cite is an AD id the spine defines, or a source file name followed by one of its headings
(`SPEC.md §4.2`, `ux-design.md "Chat screen"`). A diagram block embeds <name>.png from the diagrams
folder (the .drawio must exist, be drawn from the same spine and not be stale). A gap block is an open
point: printed where it arises and collected in the Open Points appendix, never written as fact.

Every level-1 heading of the outline must appear as a level-1 section, in order: a section with
nothing to cite keeps its heading and says so in a gap. The outline is
<design_docs_folder>/templates/<type>.md when the project has one, else the kit's
assets/design-docs/<type>.md. Word styles come from <design_docs_folder>/templates/reference.docx
when the project has one (its headers, footers and styles are kept), else built-in defaults on A4.

The document has a cover with document control, a table of contents (Word fills it on opening),
the numbered sections with each citation in grey after its block, Appendix A Open Points and
Appendix B Traceability (each AD with the sections citing it, and the ADs not cited). The stamp
(spine commit, AD hashes, source hashes) goes into the model's "stamp" so diagram-drift.py checks
documents too.

--formats confluence also writes <name>.confluence.md: the same document as Markdown for a
Confluence page (citations in brackets, open points as quotes, diagrams as images to attach), and
reports its sha256, the page title and the PNGs to attach.

Usage:
    uv run build-docx.py <project-root> <folder>/<name>.doc.json [--check] [--force] [--formats docx,confluence]

--check validates only. A .docx changed since the last build (edited or reviewed in Word) is not
replaced without --force. Prints JSON. Exit codes: 0 ok, 1 invalid model or refused, 2 error.
"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import spine as spine_mod  # noqa: E402


def _load(name: str, file: str):
    spec = importlib.util.spec_from_file_location(name, HERE / file)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


drawio = _load("build_drawio", "build-drawio.py")
drift_mod = _load("diagram_drift", "diagram-drift.py")

SKILL = HERE.parent
TYPE_NAMES = {"hld": "High-Level Design", "lld": "Low-Level Design"}
BLOCK_KINDS = ("text", "bullets", "table", "diagram", "gap")


class ModelError(Exception):
    pass


def outline_path(root: Path, kind: str) -> Path:
    own = spine_mod.design_docs_folder(root) / "templates" / f"{kind}.md"
    return own if own.is_file() else SKILL / "assets" / "design-docs" / f"{kind}.md"


def outline_sections(path: Path) -> list[str]:
    text = re.sub(r"<!--.*?-->", "", path.read_text(encoding="utf-8"), flags=re.DOTALL)
    return [m.group(1).strip() for m in re.finditer(r"^##\s+(.+)$", text, re.MULTILINE)]


def norm(heading: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", re.sub(r"^\d+(\.\d+)*\s+", "", heading.lower())).strip()


# --------------------------------------------------------------------------- validation


def validate(model: dict, root: Path) -> dict:
    """Raise ModelError listing every problem; return what the build needs."""
    problems: list[str] = []
    kind = model.get("type")
    if kind not in TYPE_NAMES:
        raise ModelError(f"type must be one of {sorted(TYPE_NAMES)}")
    if not str(model.get("title", "")).strip():
        problems.append("title is required")
    if kind == "lld" and not str(model.get("scope", "")).strip():
        problems.append("scope is required for an LLD: the component(s) it covers, or 'whole system'")

    spine_rel = str(model.get("spine", ""))
    spine_path = root / spine_rel
    if not (spine_rel.startswith(drawio.SPINE_DIR) and spine_rel.endswith("/ARCHITECTURE-SPINE.md") and spine_path.is_file()):
        raise ModelError(f"spine must be an existing {drawio.SPINE_DIR}<name>/ARCHITECTURE-SPINE.md, got {spine_rel!r}")
    sources = [spine_path]
    for rel in model.get("sources", []):
        path = (root / rel).resolve()
        try:
            inside = path.relative_to(root.resolve()).as_posix()
        except ValueError:
            inside = ""
        if not inside.startswith("_bmad-output/"):
            problems.append(f"source {rel!r} is outside _bmad-output/: design documents come from the BMad artifacts only")
        elif not path.is_file():
            problems.append(f"source not found: {rel!r}")
        else:
            sources.append(path)
    ads = spine_mod.parse_ads(spine_path.read_text(encoding="utf-8"))[0]
    names = sorted({p.name for p in sources})
    headings = {p.name: drawio.source_headings(p) for p in sources}

    def check_cites(where: str, cites: list[str]) -> None:
        if not cites:
            problems.append(f"{where}: no cite; if no source states it, make it a gap block instead")
        for c in cites:
            if drawio.AD_ID.match(c):
                if c not in ads:
                    problems.append(f"{where}: cites {c}, which the spine does not define (it defines {', '.join(ads) or 'no ADs'})")
                continue
            name = next((n for n in names if c == n or c.startswith((n + " ", n + "#", n + ":"))), None)
            if name is None:
                problems.append(f"{where}: cite {c!r} is neither an AD id nor starts with a source file name ({', '.join(names)})")
            elif not drawio.section_exists(c[len(name):], headings[name]):
                shown = "; ".join(f"{num + ' ' if num else ''}{title}" for num, title in headings[name][:12])
                problems.append(f"{where}: cite {c!r} names no section of {name} (its headings: {shown or 'none'})")

    diagrams_dir = spine_mod.diagrams_folder(root)
    diagrams: dict[str, dict] = {}
    sections = model.get("sections") or []
    if not sections:
        problems.append("sections: the document has none")
    for s_i, sec in enumerate(sections):
        where = f"section {s_i + 1} {sec.get('heading', '')!r}"
        if not str(sec.get("heading", "")).strip():
            problems.append(f"section {s_i + 1}: heading is required")
        if sec.get("level", 1) not in (1, 2, 3):
            problems.append(f"{where}: level must be 1, 2 or 3")
        if s_i == 0 and sec.get("level", 1) != 1:
            problems.append(f"{where}: the first section must be level 1")
        if not sec.get("blocks"):
            problems.append(f"{where}: no blocks; a section with nothing to say holds a gap block saying why")
        for b_i, block in enumerate(sec.get("blocks") or []):
            bw = f"{where} block {b_i + 1}"
            kinds = [k for k in BLOCK_KINDS if k in block]
            if len(kinds) != 1:
                problems.append(f"{bw}: needs exactly one of {', '.join(BLOCK_KINDS)}")
                continue
            k = kinds[0]
            if k == "gap":
                if not str(block["gap"]).strip():
                    problems.append(f"{bw}: empty gap")
            elif k == "diagram":
                name = str(block["diagram"])
                file = diagrams_dir / f"{name}.drawio"
                if not file.is_file():
                    problems.append(f"{bw}: diagram {name!r} is not in {diagrams_dir}; draw it first")
                    continue
                state = drift_mod.check(root, file)
                if state.get("spine") != spine_rel:
                    problems.append(f"{bw}: diagram {name!r} was drawn from {state.get('spine')!r}, not this document's spine")
                elif state["status"] in ("stale", "spine-missing", "unstamped"):
                    problems.append(f"{bw}: diagram {name!r} is {state['status']}; redraw it before it goes in the document")
                png = file.with_suffix(".png")
                diagrams[name] = {"title": state.get("title") or name, "png": png if png.is_file() else None,
                                  "drawn": (state.get("drawn") or "")[:10],
                                  "from": (state.get("drawn_from") or "")[:7 if "T" not in (state.get("drawn_from") or "") else 10]}
            else:
                if k == "text" and not str(block["text"]).strip():
                    problems.append(f"{bw}: empty text")
                if k == "bullets" and not (isinstance(block["bullets"], list) and all(str(b).strip() for b in block["bullets"]) and block["bullets"]):
                    problems.append(f"{bw}: bullets must be a list of non-empty strings")
                if k == "table":
                    t = block["table"]
                    cols = t.get("columns") if isinstance(t, dict) else None
                    rows = t.get("rows") if isinstance(t, dict) else None
                    if not cols or not isinstance(rows, list) or any(len(r) != len(cols) for r in rows):
                        problems.append(f"{bw}: table needs columns and rows of the same width")
                check_cites(bw, drawio.as_list(block.get("cite")))

    outline_file = outline_path(root, kind)
    wanted = outline_sections(outline_file)
    have = [norm(s.get("heading", "")) for s in sections if s.get("level", 1) == 1]
    pos = -1
    for w in wanted:
        if norm(w) not in have:
            problems.append(f"outline section {w!r} is missing: keep it, with a gap block if the sources say nothing")
            continue
        at = have.index(norm(w))
        if at < pos:
            problems.append(f"outline section {w!r} is out of order")
        pos = max(pos, at)

    if problems:
        raise ModelError("; ".join(problems))
    return {"sources": sources, "ads": ads, "diagrams": diagrams, "outline": outline_file}


# --------------------------------------------------------------------------- layout (shared by Word and Confluence)

NO_OPEN_POINTS = "None: the sources settle everything this document states."
PREFACE = ("Every statement cites the architecture decision (AD) or BMad document section it comes from, in "
           "brackets after it. What the sources do not settle is an open point, listed in Appendix A, never "
           "written as fact.")


def caption(n: int, d: dict) -> str:
    return f"Figure {n}: {d['title']}. Drawn {d['drawn']} from the spine at {d['from']}."


def uncited_line(uncited: list[str]) -> str:
    return ("ADs this document does not cite: " + ", ".join(uncited) + ".") if uncited else "Every AD in the spine is cited."


def layout(model: dict, info: dict) -> dict:
    """The document in order, numbered: headings, blocks with their cites, figures and open points."""
    counters = [0, 0, 0]
    items: list[tuple] = []
    cited: dict[str, list[str]] = {}
    open_points: list[tuple[str, str]] = []
    figures = 0
    for sec_model in model["sections"]:
        level = sec_model.get("level", 1)
        counters[level - 1] += 1
        for i in range(level, 3):
            counters[i] = 0
        number = ".".join(str(c) for c in counters[:level])
        items.append(("heading", level, f"{number} {re.sub(r'^\d+(\.\d+)*\s+', '', sec_model['heading'])}"))
        for block in sec_model["blocks"]:
            cites = drawio.as_list(block.get("cite"))
            for c in cites:
                cited.setdefault(c, [])
                if number not in cited[c]:
                    cited[c].append(number)
            if "text" in block:
                items.append(("text", str(block["text"]), cites))
            elif "bullets" in block:
                items.append(("bullets", [str(b) for b in block["bullets"]], cites))
            elif "table" in block:
                items.append(("table", [str(c) for c in block["table"]["columns"]],
                              [[str(v) for v in row] for row in block["table"]["rows"]], cites))
            elif "diagram" in block:
                figures += 1
                items.append(("figure", figures, block["diagram"], info["diagrams"][block["diagram"]]))
            elif "gap" in block:
                open_points.append((number, str(block["gap"])))
                items.append(("gap", len(open_points), str(block["gap"])))
    for g in model.get("gaps") or []:
        open_points.append(("whole document", str(g)))
    trace = [[a, info["ads"][a]["title"], ", ".join(cited[a])]
             for a in sorted((c for c in cited if drawio.AD_ID.match(c)), key=lambda a: int(a[3:]))]
    trace += [[c, "", ", ".join(cited[c])] for c in sorted(c for c in cited if not drawio.AD_ID.match(c))]
    uncited = [a for a in sorted(info["ads"], key=lambda a: int(a[3:])) if a not in cited]
    return {"items": items, "cited": cited, "open_points": open_points, "figures": figures, "trace": trace,
            "uncited": uncited}


def write_confluence(model: dict, info: dict, stamp: dict, target: Path) -> dict:
    """The same document as Markdown for a Confluence page (Atlassian MCP), diagrams as attachments to add."""
    def cell(v: str) -> str:
        return str(v).replace("|", "\\|").replace("\n", " ")

    def table(columns: list[str], rows: list[list[str]]) -> list[str]:
        return ["| " + " | ".join(cell(c) for c in columns) + " |", "|" + "---|" * len(columns),
                *["| " + " | ".join(cell(v) for v in row) + " |" for row in rows], ""]

    def cites(c: list[str]) -> str:
        return f" _[{'; '.join(c)}]_" if c else ""

    lay = layout(model, info)
    kind_name = TYPE_NAMES[model["type"]]
    built = drawio.stamp_line(stamp["_spine"], stamp["drawn"]).replace("Drawn ", "Built ")
    out = [f"> Generated by Da Vinci from the project's BMad artifacts. {built}. "
           "Change the design in the project and republish; edits made here are carried back before the next update.", "",
           *table(["Document", f"{model['title']}: {kind_name}" + (f" ({model['scope']})" if model.get("scope") else "")],
                  [["Version", model.get("version", "0.1")], ["Status", model.get("status", "Draft")],
                   ["Date", datetime.now().date().isoformat()], ["Author", model.get("author", "")]]),
           f"_{PREFACE}_", ""]
    attachments = []
    for item in lay["items"]:
        kind = item[0]
        if kind == "heading":
            out += [f"{'#' * item[1]} {item[2]}", ""]
        elif kind == "text":
            out += [item[1] + cites(item[2]), ""]
        elif kind == "bullets":
            out += [f"- {b}" + (cites(item[2]) if n == len(item[1]) - 1 else "") for n, b in enumerate(item[1])] + [""]
        elif kind == "table":
            out += table(item[1], item[2]) + [cites(item[3]).strip(), ""]
        elif kind == "figure":
            _, n, name, d = item
            if d["png"]:
                attachments.append(str(d["png"]))
                out += [f"![{caption(n, d)}]({d['png'].name})", ""]
            else:
                out += [f"**[Diagram {name}.drawio has no PNG export yet]**", ""]
            out += [f"_{caption(n, d)}_", ""]
        elif kind == "gap":
            out += [f"> **Open point {item[1]}:** {item[2]}", ""]
    out += ["# Appendix A: Open Points", ""]
    out += table(["#", "Section", "Open point"], [[str(i), s_, g] for i, (s_, g) in enumerate(lay["open_points"], 1)]) \
        if lay["open_points"] else [NO_OPEN_POINTS, ""]
    out += ["# Appendix B: Traceability", "", *table(["Source", "Title", "Cited in sections"], lay["trace"]),
            uncited_line(lay["uncited"]), ""]
    text = "\n".join(out)
    target.write_text(text, encoding="utf-8")
    return {"file": str(target), "sha256": hashlib.sha256(text.encode("utf-8")).hexdigest()[:12],
            "attachments": attachments, "title": f"{model['title']} {kind_name}" + (f": {model['scope']}" if model.get("scope") else "")}


# --------------------------------------------------------------------------- document


def _field(paragraph, instruction: str, placeholder: str = "") -> None:
    """A Word field (TOC, PAGE) in its own runs: begin, instruction, separate, placeholder, end."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def run_with(tag: str, **attrs):
        r = OxmlElement("w:r")
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        r.append(el)
        paragraph._p.append(r)
        return el

    run_with("w:fldChar", **{"w:fldCharType": "begin"})
    instr = run_with("w:instrText", **{"xml:space": "preserve"})
    instr.text = f" {instruction} "
    run_with("w:fldChar", **{"w:fldCharType": "separate"})
    if placeholder:
        paragraph.add_run(placeholder)
    run_with("w:fldChar", **{"w:fldCharType": "end"})


def _shade(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _style(doc, name: str, fallback: str = "Normal") -> str:
    try:
        doc.styles[name]
        return name
    except KeyError:
        return fallback


def write_docx(model: dict, root: Path, info: dict, stamp: dict, target: Path) -> dict:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    GREY = RGBColor(0x66, 0x66, 0x66)
    GAP = RGBColor(0xB8, 0x54, 0x50)
    reference = spine_mod.design_docs_folder(root) / "templates" / "reference.docx"
    if reference.is_file():
        doc = Document(str(reference))
        body = doc.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)
    else:
        doc = Document()
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Cm(2.2))
        doc.styles["Normal"].font.name = "Calibri"
        doc.styles["Normal"].font.size = Pt(10.5)
    sec = doc.sections[0]
    text_width = sec.page_width - sec.left_margin - sec.right_margin

    kind_name = TYPE_NAMES[model["type"]]
    subtitle = kind_name + (f": {model['scope']}" if model.get("scope") else "")
    today = datetime.now().date().isoformat()

    def cite_run(paragraph, cites: list[str]) -> None:
        run = paragraph.add_run(f"  [{'; '.join(cites)}]")
        run.font.size, run.font.color.rgb = Pt(8), GREY

    def table(columns: list[str] | None, rows: list[list], widths: list[float] | None = None):
        """A table with a shaded header row; with columns None, a key/value table with its first column bold."""
        width = len(columns) if columns else len(rows[0])
        t = doc.add_table(rows=len(rows) + (1 if columns else 0), cols=width)
        t.style = _style(doc, "Table Grid", t.style.name if t.style else "Normal")
        widths = widths or [1 / width] * width
        for r_i, values in enumerate([columns, *rows] if columns else rows):
            for c_i, value in enumerate(values):
                cell = t.cell(r_i, c_i)
                cell.width = int(text_width * widths[c_i])
                cell.text = str(value)
                header = (r_i == 0) if columns else (c_i == 0)
                if header:
                    for run in cell.paragraphs[0].runs:
                        run.font.bold = True
                    _shade(cell, "D9E2F3")
        return t

    # Cover and document control
    doc.add_paragraph(model["title"], style=_style(doc, "Title"))
    doc.add_paragraph(subtitle, style=_style(doc, "Subtitle"))
    table(None, [["Version", model.get("version", "0.1")], ["Status", model.get("status", "Draft")],
                     ["Date", today], ["Author", model.get("author", "")],
                     ["Drawn from", drawio.stamp_line(stamp["_spine"], stamp["drawn"]).replace("Drawn ", "Built ")]],
          [0.25, 0.75])
    note = doc.add_paragraph()
    r = note.add_run(PREFACE)
    r.font.size, r.font.color.rgb = Pt(9), GREY
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("Contents", style=_style(doc, "TOC Heading", _style(doc, "Heading 1")))
    _field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u', "Right-click and choose Update Field to show the contents.")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # Sections
    lay = layout(model, info)
    for item in lay["items"]:
        kind = item[0]
        if kind == "heading":
            doc.add_heading(item[2], level=item[1])
        elif kind == "text":
            cite_run(doc.add_paragraph(item[1]), item[2])
        elif kind == "bullets":
            for n, entry in enumerate(item[1]):
                p = doc.add_paragraph(entry, style=_style(doc, "List Bullet"))
                if n == len(item[1]) - 1:
                    cite_run(p, item[2])
        elif kind == "table":
            table(item[1], item[2])
            cite_run(doc.add_paragraph(), item[3])
        elif kind == "figure":
            _, n, name, d = item
            if d["png"]:
                doc.add_picture(str(d["png"]), width=text_width)
            else:
                run = doc.add_paragraph().add_run(f"[Diagram {name}.drawio: no PNG export yet. Export it from draw.io, "
                                                  "or install draw.io desktop and redraw, then rebuild this document.]")
                run.font.color.rgb = GAP
            doc.add_paragraph(caption(n, d), style=_style(doc, "Caption"))
        elif kind == "gap":
            p = doc.add_paragraph()
            run = p.add_run(f"Open point {item[1]}: ")
            run.font.bold, run.font.color.rgb = True, GAP
            p.add_run(item[2]).font.color.rgb = GAP

    # Appendices
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Appendix A: Open Points", level=1)
    if lay["open_points"]:
        table(["#", "Section", "Open point"], [[str(i), s_, g] for i, (s_, g) in enumerate(lay["open_points"], 1)],
              [0.06, 0.16, 0.78])
    else:
        doc.add_paragraph(NO_OPEN_POINTS)
    doc.add_heading("Appendix B: Traceability", level=1)
    table(["Source", "Title", "Cited in sections"], lay["trace"], [0.3, 0.45, 0.25])
    doc.add_paragraph(uncited_line(lay["uncited"]))

    # Footer with page numbers, unless the Word template brings its own
    footer = sec.footer
    if not any(p.text.strip() for p in footer.paragraphs):
        fp = footer.paragraphs[0]
        fp.text = f"{model['title']} {kind_name} v{model.get('version', '0.1')} ({model.get('status', 'Draft')})    Page "
        _field(fp, "PAGE", "1")

    # Ask Word to fill the table of contents on opening
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)

    props = doc.core_properties
    props.title, props.subject, props.author = model["title"], subtitle, model.get("author", "")
    props.comments = drawio.stamp_line(stamp["_spine"], stamp["drawn"])
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(target))
    return {"figures": lay["figures"], "open_points": len(lay["open_points"]), "cited": sorted(lay["cited"]),
            "uncited_ads": lay["uncited"]}


# --------------------------------------------------------------------------- main


def make_stamp(model: dict, root: Path, info: dict) -> dict:
    st = spine_mod.stamp(root, info["sources"][0])
    drawn = datetime.now(timezone.utc).isoformat(timespec="seconds")
    cited = sorted({c for s in model["sections"] for b in s["blocks"] for c in drawio.as_list(b.get("cite"))
                    if drawio.AD_ID.match(c)}, key=lambda a: int(a[3:]))
    return {
        "_spine": st, "drawn": drawn, "arch_diagram": "1", "arch_type": model["type"],
        "arch_title": f"{model['title']} {TYPE_NAMES[model['type']]}" + (f": {model['scope']}" if model.get("scope") else ""),
        "spine": st["spine"], "spine_commit": st["spine_commit"], "spine_committed": st["spine_committed"],
        "spine_dirty": str(st["spine_dirty"]).lower(), "spine_modified": st["spine_modified"],
        "ad_hashes": json.dumps(st["ad_hashes"], separators=(",", ":")), "cited_ads": ",".join(cited),
        "source_hashes": json.dumps({p.resolve().relative_to(root.resolve()).as_posix(): spine_mod.file_hash(p)
                                     for p in info["sources"][1:]}, separators=(",", ":")),
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("project_root")
    parser.add_argument("model", help="<folder>/<name>.doc.json; <name>.docx is written beside it")
    parser.add_argument("--check", action="store_true", help="Validate only")
    parser.add_argument("--force", action="store_true", help="Replace a .docx changed since the last build")
    parser.add_argument("--formats", default="docx", help="Comma-separated: docx, confluence (default: docx)")
    args = parser.parse_args(argv)
    root = Path(args.project_root).resolve()
    model_path = Path(args.model)
    if not model_path.is_absolute() and not model_path.exists():
        model_path = root / model_path
    if not model_path.name.endswith(".doc.json"):
        print(json.dumps({"status": "error", "message": "the model file must be named <name>.doc.json"}))
        return 1
    target = model_path.with_name(model_path.name[: -len(".doc.json")] + ".docx")
    try:
        model = json.loads(model_path.read_text(encoding="utf-8"))
        info = validate(model, root)
    except ModelError as exc:
        print(json.dumps({"status": "error", "message": str(exc)}))
        return 1
    except (OSError, json.JSONDecodeError) as exc:
        print(json.dumps({"status": "error", "message": str(exc)}))
        return 2
    if args.check:
        print(json.dumps({"status": "valid", "outline": str(info["outline"]),
                          "diagrams": {k: bool(v["png"]) for k, v in info["diagrams"].items()}}, indent=2))
        return 0
    formats = {f.strip().lower() for f in args.formats.split(",") if f.strip()}
    if not formats or formats - {"docx", "confluence"}:
        print(json.dumps({"status": "error", "message": "--formats takes docx, confluence or both"}))
        return 1
    saved = (model.get("stamp") or {}).get("docx_mtime")
    if "docx" in formats and target.exists() and not args.force and (saved is None or abs(target.stat().st_mtime - saved) > 1):
        print(json.dumps({"status": "error", "message": f"{target.name} changed since the last build (edited or reviewed in "
                          "Word?); carry those changes into the model first, or pass --force to replace it"}))
        return 1
    stamp = make_stamp(model, root, info)
    report: dict = {}
    if "docx" in formats:
        try:
            report = write_docx(model, root, info, stamp, target)
        except PermissionError:
            print(json.dumps({"status": "error", "message": f"cannot write {target}: is it open in Word?"}))
            return 2
    if "confluence" in formats:
        lay = layout(model, info)
        report = report or {"figures": lay["figures"], "open_points": len(lay["open_points"]),
                            "cited": sorted(lay["cited"]), "uncited_ads": lay["uncited"]}
        report["confluence"] = write_confluence(model, info, stamp, target.with_suffix(".confluence.md"))
    stamp.pop("_spine")
    old = model.get("stamp") or {}
    model["stamp"] = {**stamp, "docx_mtime": target.stat().st_mtime if "docx" in formats else old.get("docx_mtime")}
    model_path.write_text(json.dumps(model, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    missing_png = [k for k, v in info["diagrams"].items() if not v["png"]]
    print(json.dumps({"status": "success", "file": str(target) if "docx" in formats else None, "outline": str(info["outline"]),
                      "stamp": drawio.stamp_line(spine_mod.stamp(root, info["sources"][0]), stamp["drawn"]).replace("Drawn ", "Built "),
                      **report, "diagrams_without_png": missing_png}, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
